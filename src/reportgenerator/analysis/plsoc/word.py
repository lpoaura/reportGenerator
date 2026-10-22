"""Insertion dans le Word du rapport PLSOC : tableaux et suites de blocs (secteurs, fiches de mesures).

Un placeholder de tableau ou de blocs doit être seul dans son paragraphe ; ce paragraphe est remplacé.
Blocs : ("titre", niveau, texte) | ("texte", texte) | ("a_rediger", texte) | ("legende", texte)
        | ("tableau", lignes, colonnes[, largeurs_mm]) | ("image", chemin[, largeur_mm]).
"""

from pathlib import Path

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_COLOR_INDEX, WD_PARAGRAPH_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from docx.text.paragraph import Paragraph

from reportgenerator.analysis.common.tables.styles import LPO_BLUE
from reportgenerator.analysis.common.tables.utils import set_cell_background
from reportgenerator.core.renderer import compact_image

TOTAL_FOND = "E7E6E6"


def _paragraphe(document, placeholder):
    return next((p for p in document.paragraphs if placeholder in p.text), None)


def _style(document, *noms):
    existants = {s.name for s in document.styles}
    return next((n for n in noms if n in existants), None)


def _entete_repete(row):
    tr_pr = row._tr.get_or_add_trPr()
    e = OxmlElement("w:tblHeader")
    e.set(qn("w:val"), "true")
    tr_pr.append(e)


def _texte_cellule(cell, texte, taille, gras=False, blanc=False, centre=False):
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    # pas de texte justifié dans les cellules (style Normal du template) : espaces étirés sur les colonnes étroites
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER if centre else WD_PARAGRAPH_ALIGNMENT.LEFT
    run = p.add_run("" if texte is None else str(texte))
    run.font.size = Pt(taille)
    run.bold = gras
    if blanc:
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


def construire_tableau(document, lignes, colonnes, largeurs_mm=None, taille=8.5):
    """colonnes = [(clé, libellé)] ; une ligne avec total=True est en gras sur fond gris."""
    t = document.add_table(rows=1, cols=len(colonnes))
    style = _style(document, "Table Grid")
    if style:
        t.style = style
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (_, lib) in enumerate(colonnes):
        c = t.rows[0].cells[i]
        set_cell_background(c, LPO_BLUE)
        _texte_cellule(c, lib, taille, gras=True, blanc=True, centre=True)
    _entete_repete(t.rows[0])
    for ligne in lignes:
        cells = t.add_row().cells
        for i, (cle, _) in enumerate(colonnes):
            _texte_cellule(cells[i], ligne.get(cle, ""), taille, gras=bool(ligne.get("total")))
            if ligne.get("total"):
                set_cell_background(cells[i], TOTAL_FOND)
    if largeurs_mm:
        for row in t.rows:
            for i, w in enumerate(largeurs_mm):
                row.cells[i].width = Mm(w)
    return t


def inserer_tableau(document, placeholder, data, colonnes, largeurs_mm=None, vide="Aucune donnée.", taille=8.5):
    p = _paragraphe(document, placeholder)
    if p is None:
        return
    if not data:
        for r in p.runs:
            r.text = ""
        p.add_run(vide).italic = True
        return
    t = construire_tableau(document, data, colonnes, largeurs_mm, taille)
    p._p.addnext(t._tbl)
    p._p.getparent().remove(p._p)


def _nouveau_paragraphe(apres, parent):
    e = OxmlElement("w:p")
    apres.addnext(e)
    return Paragraph(e, parent)


def inserer_blocs(document, placeholder, data):
    p = _paragraphe(document, placeholder)
    if p is None:
        return
    parent, ancre = p._parent, p._p
    for bloc in data:
        genre = bloc[0]
        if genre == "tableau":
            if not bloc[1]:
                q = _nouveau_paragraphe(ancre, parent)
                q.add_run("Aucune donnée.").italic = True
                ancre = q._p
                continue
            t = construire_tableau(document, bloc[1], bloc[2], bloc[3] if len(bloc) > 3 else None)
            ancre.addnext(t._tbl)
            ancre = t._tbl
            continue
        q = _nouveau_paragraphe(ancre, parent)
        ancre = q._p
        if genre == "titre":
            style = _style(document, f"Heading {bloc[1]}", f"Titre {bloc[1]}")
            if style:
                q.style = document.styles[style]
            q.add_run(bloc[2])
        elif genre == "texte":
            q.add_run(bloc[1])
        elif genre == "a_rediger":
            r = q.add_run(bloc[1])
            r.italic = True
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        elif genre == "legende":
            style = _style(document, "Caption", "Légende")
            if style:
                q.style = document.styles[style]
            q.add_run(bloc[1])
        elif genre == "image":
            chemin = Path(bloc[1])
            if chemin.exists():
                q.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
                q.add_run().add_picture(str(compact_image(chemin)), width=Mm(bloc[2] if len(bloc) > 2 else 165))
    p._p.getparent().remove(p._p)
