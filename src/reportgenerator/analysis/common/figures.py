"""Insertion d'une liste de figures (cartes d'atlas) à la place d'un placeholder du Word.

    {{ATLAS_ESPECES}}  ->  [titre du groupe (style Titre 3)]
                           image (réduite pour tenir dans la page)
                           Figure n : légende (style Légende, numérotation Word)
                           [texte à rédiger, surligné en jaune]
                           ... figure suivante
"""

from pathlib import Path

from docx.enum.text import WD_COLOR_INDEX, WD_PARAGRAPH_ALIGNMENT
from docx.image.image import Image as DocxImage
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm
from docx.text.paragraph import Paragraph

from reportgenerator.core.renderer import compact_image

# taille maximale d'une carte : la légende et le texte à rédiger tiennent sur la même page (A4, marges LPO)
MAX_WIDTH, MAX_HEIGHT = Mm(170), Mm(195)


def _new_paragraph_after(element, parent):
    p = OxmlElement("w:p")
    element.addnext(p)
    return Paragraph(p, parent)


def _style(document, *names):
    """Premier style existant parmi names (les templates LPO sont en anglais ou en français)."""
    existing = {s.name for s in document.styles}
    return next((n for n in names if n in existing), None)


def _figures_avant(document, element) -> int:
    """Nombre de champs « SEQ Figure » avant element (numéro affiché sans mise à jour des champs)."""
    n = 0
    for el in document.element.body.iter():
        if el is element:
            break
        if el.tag == qn("w:fldSimple") and "SEQ Figure" in (el.get(qn("w:instr")) or ""):
            n += 1
        elif el.tag == qn("w:instrText") and "SEQ Figure" in (el.text or ""):
            n += 1
    return n


def _caption(paragraph, text, numero=None):
    """« Figure <SEQ> : texte » : le numéro est un champ Word (mis à jour avec F9, comme le sommaire) ; numero =
    valeur affichée en attendant (sinon le champ reste vide tant qu'il n'est pas mis à jour)."""
    paragraph.add_run("Figure ")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), r" SEQ Figure \* ARABIC ")
    if numero is not None:
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = str(numero)
        r.append(t)
        fld.append(r)
    paragraph._p.append(fld)
    paragraph.add_run(f" : {text}")


def insert_figure_list(document, placeholder, data, note="[Photo et texte de l'espèce à rédiger]", niveau=3,
                       hauteur_max_mm=None):
    """data = [{"titre": str | None, "avertissement": str | None (optionnel),
                "figures": [{"path": Path, "legende": str}, ...]}, ...].
    Les figures absentes (image non produite) sont ignorées ; un groupe sans figure n'a pas de titre.
    niveau : niveau de titre des groupes (3 = Titre 3) ; hauteur_max_mm : hauteur maximale d'une image (195 mm par
    défaut, pour que la légende et le texte à rédiger tiennent sur la même page)."""
    max_height = Mm(hauteur_max_mm) if hauteur_max_mm else MAX_HEIGHT
    paragraph = next((p for p in document.paragraphs if placeholder in p.text), None)
    if paragraph is None:
        return
    heading = _style(document, f"Heading {niveau}", f"Titre {niveau}")
    caption = _style(document, "Caption", "Légende")
    anchor, parent = paragraph._p, paragraph._parent
    numero = _figures_avant(document, anchor)
    for bloc in data:
        figures = [f for f in bloc["figures"] if Path(f["path"]).exists()]
        if not figures:
            continue
        if bloc.get("titre"):
            p = _new_paragraph_after(anchor, parent)
            p.add_run(bloc["titre"])
            if heading:
                p.style = heading
            p.paragraph_format.keep_with_next = True   # le titre reste avec sa première carte
            anchor = p._p
        if bloc.get("avertissement"):   # mise en garde de lecture, en italique, avant les cartes du groupe
            p = _new_paragraph_after(anchor, parent)
            p.add_run(bloc["avertissement"]).italic = True
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.keep_together = True   # pas coupé entre deux pages
            anchor = p._p
        for fig in figures:
            img = _new_paragraph_after(anchor, parent)
            img.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            chemin = compact_image(fig["path"])  # JPEG pour le Word (le PNG reste dans maps/)
            image = DocxImage.from_file(str(chemin))
            ratio = image.px_height / image.px_width
            img.add_run().add_picture(str(chemin), width=min(MAX_WIDTH, int(max_height / ratio)))
            img.paragraph_format.keep_with_next = True     # la carte reste avec sa légende
            cap = _new_paragraph_after(img._p, parent)
            numero += 1
            _caption(cap, fig["legende"], numero)
            if caption:
                cap.style = caption
            anchor = cap._p
            if note:
                txt = _new_paragraph_after(anchor, parent)
                txt.add_run(note).font.highlight_color = WD_COLOR_INDEX.YELLOW
                anchor = txt._p
    # le paragraphe du placeholder disparaît
    paragraph._p.getparent().remove(paragraph._p)
