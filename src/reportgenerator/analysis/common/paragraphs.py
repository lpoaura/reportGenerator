"""Insertion d'une liste de paragraphes (un par élément) à la place d'un placeholder du Word.

    {{LISTE}}  ->  « Nom en gras » + texte      (mise en forme du paragraphe du placeholder)
                   « Nom en gras » + texte
                   ...
"""

import copy

from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

VIDE = "Aucune espèce concernée."


def insert_paragraph_list(document, placeholder, data, vide=VIDE):
    """data = [{"nom": texte en gras, "texte": suite}, ...] ; liste vide : un paragraphe `vide` en italique."""
    paragraph = next((p for p in document.paragraphs if placeholder in p.text), None)
    if paragraph is None:
        return
    anchor = paragraph._p
    for item in data or [{"nom": "", "texte": vide, "italique": True}]:
        p_el = copy.deepcopy(paragraph._p)
        for child in list(p_el):
            if child.tag != qn("w:pPr"):
                p_el.remove(child)
        anchor.addnext(p_el)
        p = Paragraph(p_el, paragraph._parent)
        if item.get("nom"):
            p.add_run(item["nom"]).bold = True
        run = p.add_run(item.get("texte", ""))
        run.italic = bool(item.get("italique"))
        anchor = p_el
    paragraph._p.getparent().remove(paragraph._p)
