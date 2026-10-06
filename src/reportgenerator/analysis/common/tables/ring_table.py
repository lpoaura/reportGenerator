"""
Tableaux Word "par anneau" (dossier éolien) :

    | Nom vern. | Nom sci. | ... | Ensemble des zones  | Zone d'étude        | Contour 1 km ...
    |           |          |     | Nb données | Dern.  | Nb données | Dern.  | ...

Reprend la mise en forme de l'ancien script mise_en_forme_test.py (double
en-tête fusionné, couleurs liste rouge, noms scientifiques en italique),
directement dans le Word au lieu d'un Excel à recopier.
"""

from decimal import Decimal

from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from reportgenerator.analysis.common.tables.utils import (
    apply_status_style_species, italicize_cell_species, set_cell_background,
    set_cell_font, set_cell_text_color)

LPO_BLUE = "0088CC"
WHITE = "FFFFFF"
EMPTY_TEXT = "Aucune espèce concernée."
FONT = "LPO Light"  # police du corps de texte des templates LPO


def _fmt_value(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (float, Decimal)) and value == int(value):
        return str(int(value))
    return str(value)


def _find_paragraph(document, placeholder):
    for paragraph in document.paragraphs:
        if placeholder in paragraph.text:
            return paragraph
    return None


def _repeat_as_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


# ordre imposé par le schéma OOXML des enfants de w:tblPr (Word refuse un ordre différent)
TBLPR_ORDER = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize",
               "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook",
               "tblCaption", "tblDescription"]


def _set_tblpr_child(tbl_pr, name, **attrs):
    """Remplace / insère w:<name> dans tblPr à la place prévue par le schéma."""
    for el in tbl_pr.findall(qn(f"w:{name}")):
        tbl_pr.remove(el)
    el = OxmlElement(f"w:{name}")
    for key, value in attrs.items():
        el.set(qn(f"w:{key}"), value)
    later = TBLPR_ORDER[TBLPR_ORDER.index(name) + 1:]
    following = next((c for c in tbl_pr if c.tag.split("}")[1] in later), None)
    if following is not None:
        following.addprevious(el)
    else:
        tbl_pr.append(el)


def _fit_page_width(table, n_cols, widths=None):
    """Largeur = 100 % de la page, colonnes ajustées au contenu (proportions `widths` si fournies).
    (python-docx fixe sinon des colonnes à la largeur de la dernière section,
    ce qui peut déborder de la page.)"""
    tbl_pr = table._tbl.tblPr
    _set_tblpr_child(tbl_pr, "tblW", w="5000", type="pct")
    _set_tblpr_child(tbl_pr, "tblLayout", type="autofit")
    weights = widths or [1] * n_cols
    for grid_col, weight in zip(table._tbl.tblGrid.findall(qn("w:gridCol")), weights):
        grid_col.set(qn("w:w"), str(int(9000 * weight / sum(weights))))


def _style_header(table, n_rows, font_size):
    # les cellules fusionnées apparaissent plusieurs fois : on ne les style qu'une fois.
    # (on garde les éléments pour comparer par identité ; id() d'un proxy lxml
    # temporaire peut être réutilisé et faire sauter des cellules)
    seen = []
    for row in table.rows[:n_rows]:
        _repeat_as_header(row)
        for cell in row.cells:
            if any(cell._tc is tc for tc in seen):
                continue
            seen.append(cell._tc)
            set_cell_background(cell, LPO_BLUE)
            set_cell_text_color(cell, WHITE)
            set_cell_font(cell, bold=True, size=font_size)
            for p in cell.paragraphs:
                p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER


def _fill_rows(table, data, keys, status_fields, italic_fields, font_size):
    if not data:
        row = table.add_row()
        cell = row.cells[0].merge(row.cells[-1])
        cell.text = EMPTY_TEXT
        set_cell_font(cell, italic=True, size=font_size, font_name=FONT)
        return

    for item in data:
        cells = table.add_row().cells
        for idx, key in enumerate(keys):
            cell = cells[idx]
            value = _fmt_value(item.get(key))
            cell.text = value
            set_cell_font(cell, size=font_size, font_name=FONT)
            if idx >= 2:
                cell.paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            if key in italic_fields:
                italicize_cell_species(cell)
            if key in status_fields:
                apply_status_style_species(cell, value)


def _place(paragraph, placeholder, table):
    paragraph.text = paragraph.text.replace(placeholder, "")
    paragraph._element.addnext(table._element)


def insert_lpo_table(document, placeholder, data, columns, status_fields=(), italic_fields=(), font_size=9,
                     widths=None):
    """Tableau simple : columns = [(clé, libellé), ...] ; widths = largeurs relatives des colonnes."""
    paragraph = _find_paragraph(document, placeholder)
    if paragraph is None:
        return
    table = document.add_table(rows=1, cols=len(columns))
    _fit_page_width(table, len(columns), widths)
    for idx, (_, label) in enumerate(columns):
        table.cell(0, idx).text = label
    _style_header(table, 1, font_size)
    _fill_rows(table, data, [k for k, _ in columns], status_fields, italic_fields, font_size)
    _place(paragraph, placeholder, table)


def insert_ring_table(
    document,
    placeholder,
    data,
    fixed_columns,
    groups,
    metrics,
    status_fields=(),
    italic_fields=("lb_nom",),
    font_size=8,
):
    """Tableau à double en-tête.

    - fixed_columns : [(clé, libellé)] fusionnées sur les 2 lignes d'en-tête
    - groups        : [(suffixe, libellé)] ex. [("all", "Ensemble des zones"), (0, "Zone d'étude")]
    - metrics       : [(clé, libellé)] ex. [("nb_obs", "Nb données"), ("derniere_annee", "Dernière obs.")]
      la valeur lue dans chaque ligne est item[f"{clé}__{suffixe}"].
    """
    paragraph = _find_paragraph(document, placeholder)
    if paragraph is None:
        return

    n_cols = len(fixed_columns) + len(groups) * len(metrics)
    table = document.add_table(rows=2, cols=n_cols)
    _fit_page_width(table, n_cols)

    for idx, (_, label) in enumerate(fixed_columns):
        table.cell(0, idx).merge(table.cell(1, idx)).text = label

    col = len(fixed_columns)
    for _, group_label in groups:
        top = table.cell(0, col)
        if len(metrics) > 1:
            top = top.merge(table.cell(0, col + len(metrics) - 1))
        top.text = group_label
        for j, (_, metric_label) in enumerate(metrics):
            table.cell(1, col + j).text = metric_label
        col += len(metrics)

    _style_header(table, 2, font_size)

    keys = [k for k, _ in fixed_columns] + [
        f"{metric}__{suffix}" for suffix, _ in groups for metric, _ in metrics
    ]
    _fill_rows(table, data, keys, status_fields, italic_fields, font_size)
    _place(paragraph, placeholder, table)
