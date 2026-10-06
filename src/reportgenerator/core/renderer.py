"""
renderer.py

Remplit un template Word à partir des AnalysisResult des briques.
Ne connaît rien du métier : il remplace des placeholders {{CLE}} par
  - du texte   (result.texts)
  - un tableau (result.tables)
  - une image  (result.images ; {{cle}} ou {{cle.png}} acceptés)

Les placeholders sont cherchés dans tout le document : corps, cellules de
tableaux, zones de texte (page de garde), en-têtes et pieds de page.
La mise en forme du texte autour du placeholder est conservée.
"""

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.image.image import Image as DocxImage
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm
from docx.text.paragraph import Paragraph

LAYOUTS = {
    "normal": {"width": Inches(6), "height": None, "full_page": False},
    "A4": {"width": Mm(170), "height": Mm(247), "full_page": True},
    "A3": {"width": Mm(250), "height": Mm(370), "full_page": True},
    # image placée dans le paragraphe du placeholder, réduite pour tenir dans la page
    # (proportions conservées) ; le template gère lui-même les sauts de page
    "page": {"width": Mm(170), "height": Mm(225), "full_page": False, "fit": True},
}

PLACEHOLDER_RE = re.compile(r"\{\{\s*([^{}]+?)\s*\}\}")


# =========================================================
# PARCOURS DU DOCUMENT
# =========================================================


def iter_paragraphs(document, include_headers=True):
    """Tous les paragraphes : corps, tableaux (imbriqués), zones de texte,
    puis en-têtes / pieds de page."""
    body = document._body
    for p in document.element.body.iter(qn("w:p")):
        yield Paragraph(p, body)

    if not include_headers:
        return

    seen = set()
    for section in document.sections:
        for part in (
            section.header, section.footer,
            section.first_page_header, section.first_page_footer,
            section.even_page_header, section.even_page_footer,
        ):
            if part.is_linked_to_previous or id(part._element) in seen:
                continue
            seen.add(id(part._element))
            for p in part._element.iter(qn("w:p")):
                yield Paragraph(p, part)


def _runs_text(paragraph):
    return "".join(run.text for run in paragraph.runs)


def replace_in_paragraph(paragraph, placeholder: str, value: str) -> int:
    """Remplace `placeholder` par `value` en conservant la mise en forme.

    Word découpe souvent un placeholder sur plusieurs "runs" ; le texte de
    remplacement prend le style du premier run concerné, les suivants sont
    vidés. Retourne le nombre de remplacements.
    """
    count = 0
    search_from = 0
    while True:
        runs = paragraph.runs
        full = "".join(run.text for run in runs)
        start = full.find(placeholder, search_from)
        if start < 0:
            return count
        end = start + len(placeholder)

        pos = 0
        first = True
        for run in runs:
            run_start, run_end = pos, pos + len(run.text)
            pos = run_end
            if run_end <= start or run_start >= end:
                continue
            before = run.text[: max(start - run_start, 0)]
            after = run.text[end - run_start:] if end < run_end else ""
            if first:
                run.text = before + value + after
                first = False
            else:
                run.text = after

        count += 1
        search_from = start + len(value)


def find_placeholders(document) -> set[str]:
    """Clés des placeholders, hors balises de section ({{#nom}} / {{/nom}})."""
    keys = set()
    for paragraph in iter_paragraphs(document):
        keys.update(k for k in PLACEHOLDER_RE.findall(_runs_text(paragraph)) if k[0] not in "#/")
    return keys


SECTION_RE = re.compile(r"^\{\{\s*([#/])\s*([^{}]+?)\s*\}\}$")


def apply_sections(document, keep: set[str] | None) -> list[str]:
    """Sections conditionnelles : paragraphes {{#nom}} ... {{/nom}} (seuls dans leur paragraphe,
    dans le corps du document). Si nom est dans `keep` (ou keep est None), seules les balises
    sont retirées ; sinon toute la section est supprimée. Retourne les sections supprimées."""
    body = document.element.body

    def marker(el):
        if el.tag != qn("w:p"):
            return None
        m = SECTION_RE.match(_runs_text(Paragraph(el, document._body)).strip())
        return m.groups() if m else None

    removed = []
    while True:
        children = list(body.iterchildren())
        markers = [marker(el) for el in children]
        i = next((k for k, m in enumerate(markers) if m and m[0] == "#"), None)
        if i is None:
            return removed
        name = markers[i][1]
        end = next((k for k in range(i + 1, len(children)) if markers[k] == ("/", name)), None)
        if end is None:
            raise ValueError(f"Section {{{{#{name}}}}} sans balise de fin {{{{/{name}}}}}")
        if keep is None or name in keep:
            to_remove = [children[i], children[end]]
        else:
            to_remove = children[i:end + 1]
            removed.append(name)
        for el in to_remove:
            body.remove(el)


def template_placeholders(template_path: Path) -> set[str]:
    """Clés des placeholders d'un template (sans extension .png)."""
    return {normalize_key(k) for k in find_placeholders(Document(template_path))}


def normalize_key(key: str) -> str:
    return key[:-4] if key.lower().endswith(".png") else key


# =========================================================
# INSERTIONS
# =========================================================


def insert_paragraph_after(paragraph):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    return Paragraph(new_p, paragraph._parent)


def _insert_image(paragraph, placeholder, image_path, layout):
    config = LAYOUTS.get(layout)
    if not config:
        raise ValueError(f"Layout inconnu : {layout}")

    replace_in_paragraph(paragraph, placeholder, "")

    if config.get("fit"):
        image = DocxImage.from_file(str(image_path))
        ratio = image.px_height / image.px_width
        width = min(config["width"], int(config["height"] / ratio))
        paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        paragraph.add_run().add_picture(str(image_path), width=width)
        return

    if not config["full_page"]:
        paragraph.add_run().add_picture(str(image_path), width=config["width"])
        return

    image_p = insert_paragraph_after(paragraph)
    image_p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    image_p.add_run().add_picture(
        str(image_path), width=config["width"], height=config["height"]
    )
    page_break_p = insert_paragraph_after(image_p)
    page_break_p.add_run().add_break()


def dedupe_drawing_ids(document) -> int:
    """Rend uniques les identifiants de dessin (wp:docPr/@id) sur tout le document.

    python-docx choisit l'id d'une nouvelle image en ne regardant que le corps : il peut
    reprendre un id déjà utilisé dans un en-tête ou pied de page (logo...). Word juge
    alors le fichier endommagé. Retourne le nombre d'ids renumérotés."""
    tag = qn("wp:docPr")
    used = set()
    for part in document.part.package.iter_parts():
        element = getattr(part, "_element", None)
        if part is document.part or element is None:
            continue
        used.update(int(el.get("id")) for el in element.iter(tag) if el.get("id", "").isdigit())
    body_ids = [el for el in document.element.iter(tag)]
    next_id = max(used | {int(el.get("id")) for el in body_ids if el.get("id", "").isdigit()}, default=0) + 1
    seen, changed = set(), 0
    for el in body_ids:
        current = int(el.get("id")) if el.get("id", "").isdigit() else None
        if current is None or current in used or current in seen:
            el.set("id", str(next_id))
            current = next_id
            next_id += 1
            changed += 1
        seen.add(current)
    return changed


def _replace_images(document, images):
    for key, block in images.items():
        path = Path(block.path)
        if not path.exists():
            print(f"[AVERTISSEMENT] Image absente pour {{{{{key}}}}} : {path}")
            continue
        for placeholder in ("{{" + key + "}}", "{{" + key + ".png}}"):
            for paragraph in list(iter_paragraphs(document, include_headers=False)):
                if placeholder in _runs_text(paragraph):
                    _insert_image(paragraph, placeholder, path, block.layout)


# =========================================================
# RENDU
# =========================================================


def merge_results(results) -> tuple[dict, dict, dict]:
    """Fusionne texts / tables / images de toutes les briques.
    Deux briques ne peuvent pas fournir la même clé."""
    texts, tables, images = {}, {}, {}
    owners = {}
    for name, result in results.items():
        for target, source in ((texts, result.texts), (tables, result.tables), (images, result.images)):
            for key, value in source.items():
                if key in owners:
                    raise ValueError(
                        f"Placeholder {{{{{key}}}}} fourni par deux briques : "
                        f"{owners[key]} et {name}"
                    )
                owners[key] = name
                target[key] = value
    return texts, tables, images


def render_report(template_path: Path, results: dict, output_file: Path, sections: set[str] | None = None) -> set[str]:
    """Génère le Word. Retourne les placeholders restés sans valeur.
    sections : briques exécutées (sections conditionnelles {{#brique}} gardées) ; None = tout garder."""
    texts, tables, images = merge_results(results)
    document = Document(template_path)

    # 0. sections conditionnelles
    for name in apply_sections(document, sections):
        print(f"Section '{name}' retirée (analyse non demandée)")

    # 1. textes
    for paragraph in iter_paragraphs(document):
        if "{{" not in _runs_text(paragraph):
            continue
        for key, value in texts.items():
            replace_in_paragraph(paragraph, "{{" + key + "}}", "" if value is None else str(value))

    # 2. tableaux (fonctions d'insertion existantes)
    for key, block in tables.items():
        block.insert(document=document, placeholder="{{" + key + "}}", data=block.data, **block.options)

    # 3. images
    _replace_images(document, images)
    dedupe_drawing_ids(document)

    unresolved = find_placeholders(document)
    for key in sorted(unresolved):
        print(f"[AVERTISSEMENT] Placeholder sans valeur : {{{{{key}}}}}")

    document.save(output_file)
    return unresolved
