"""Tests du dossier PLSOC Œdicnème (sans base : données factices)."""

from datetime import date

import pytest
from docx import Document

from reportgenerator.analysis.plsoc import calculs as k
from reportgenerator.analysis.plsoc.deductions import deduire
from reportgenerator.analysis.plsoc.queries import PlsocQueries
from reportgenerator.analysis.plsoc.word import inserer_blocs, inserer_tableau
from reportgenerator.core.context import ReportContext
from reportgenerator.core.dossier import load_dossier, resolve_dossier
from reportgenerator.core.pipeline import check_dossier, run_pipeline
from reportgenerator.core.renderer import find_placeholders

NOMENC = {
    "ref_succes_repro": {
        c: {"code": c, "libelle": lib}
        for c, lib in (
            ("succes_certain", "Succès certain"),
            ("succes_probable", "Succès probable"),
            ("en_cours", "En cours"),
            ("indetermine", "Indéterminé"),
            ("echec_probable", "Échec probable"),
            ("echec_certain", "Échec certain"),
        )
    },
    "ref_stade_repro": {
        "presence_de_poussins": {"libelle": "Présence de poussins"},
        "couvaison_observee": {"libelle": "Couvaison observée"},
    },
    "ref_milieu": {
        "agricole": {"libelle": "Agricole"},
        "bassin": {"libelle": "Bassin"},
        "parcelle_compensatoire": {"libelle": "Parcelle compensatoire"},
    },
    "ref_cause_echec": {"predation": {"libelle": "Prédation"}, "travaux_agricoles": {"libelle": "Travaux agricoles"}},
    "ref_degre_menace": {},
    "ref_type_protection": {"piquetage": {"libelle": "Piquetage"}},
    "ref_statut_evenement": {"COUV": {"libelle": "Couvaison observée"}},
}
ZONE = "MULTIPOLYGON(((840000 6510000,860000 6510000,860000 6530000,840000 6530000,840000 6510000)))"
SECTEUR = "MULTIPOLYGON(((840000 6510000,850000 6510000,850000 6530000,840000 6530000,840000 6510000)))"
MESURE = "MULTIPOLYGON(((845000 6515000,845100 6515000,845100 6515100,845000 6515100,845000 6515000)))"


def couple(cle, annee, **kw):
    return {
        "source": "mergin",
        "cle_couple": cle,
        "code_couple": cle,
        "annee": annee,
        "date_premiere_obs": date(annee, 3, 10),
        "date_derniere_obs": date(annee, 6, 30),
        "lieu_dit": None,
        "milieu": None,
        "culture": None,
        "statut_nidification": "Nicheur certain",
        "nb_nids": 1,
        "nb_nids_succes": 1,
        "nb_nids_proteges": 0,
        "nb_jeunes_envol": 1,
        "succes_repro": "succes_certain",
        "stade_repro": "presence_de_poussins",
        "degre_menace": "non_menace",
        "code_insee": "69044",
        "nom_commune": "Chassieu",
        "nom_epci": "Métropole de Lyon",
        "nom_secteur": "Chassieu-Bron",
        "id_mesure_compensatoire": None,
        "zonage_plu": None,
        "x": 845050.0,
        "y": 6515050.0,
        **kw,
    }


def nid(cle, cle_couple, annee, succes, jours, **kw):
    return {
        "source": "mergin",
        "cle_nid": cle,
        "cle_couple": cle_couple,
        "code_nid": cle,
        "annee": annee,
        "numero_nichee": 1,
        "stade_repro": "couvaison_observee",
        "succes_repro": succes,
        "succes_source": "validé",
        "nb_oeufs": 2,
        "nb_poussins": 1,
        "nb_jeunes_envol": 1,
        "type_protection": None,
        "protege": False,
        "date_couvaison": date(annee, 4, 1),
        "date_eclosion": None,
        "date_echec": None,
        "date_protection": None,
        "mayfield_jours_exposition": jours,
        "mayfield_source": "validé",
        "cause_fin": None,
        "milieu": None,
        "agricole": False,
        "code_insee": "69044",
        "nom_commune": "Chassieu",
        "nom_epci": "Métropole de Lyon",
        "nom_secteur": "Chassieu-Bron",
        "id_mesure_compensatoire": None,
        "x": 845050.0,
        "y": 6515050.0,
        **kw,
    }


# --------------------------------------------------------------------------------------------------- calculs
def test_mayfield_methode_des_rapports():
    # rapport 2024 : 32 échecs certains, 1866,5 jours → survie 0,98286, probabilité 33,06 % sur 64 jours
    nids = [nid(f"n{i}", "c", 2024, "echec_certain", 0) for i in range(32)]
    nids += [nid("e", "c", 2024, "succes_certain", 1866.5)]
    m = k.mayfield(nids, 2024, 64)
    assert m["echecs"] == 32 and m["jours"] == 1866.5
    assert round(m["survie"], 5) == 0.98286 and round(m["probabilite"], 4) == 0.3306
    assert "0,98286" in k.texte_mayfield(m, 2024)


def test_bilan_et_succes():
    cs = [couple("a", 2024), couple("b", 2024, nb_nids=0)]
    ns = [
        nid("n1", "a", 2024, "succes_certain", 10),
        nid("n2", "a", 2024, "echec_probable", 5, succes_source="calculé"),
    ]
    b = k.bilan_annee(cs, ns, 2024)
    assert (b["couples"], b["nicheurs"], b["nids"], b["taux_succes"], b["valides"]) == (2, 1, 2, 50, 1)
    s = k.succes(ns, 2024, NOMENC)
    assert s["non_valides"] == 1 and "provisoire" in s["texte"]


def test_tableau_communes_sous_totaux():
    cs = [
        couple("a", 2024),
        couple("b", 2024, nom_commune="Bron"),
        couple("c", 2024, nom_epci="CCEL", nom_commune="Genas"),
    ]
    rows = k.tableau_communes(cs, [], 2024)
    totaux = [r for r in rows if r.get("total")]
    assert len(totaux) == 1 and totaux[0]["couples"] == 2 and rows[0]["collectivite"] == "Métropole de Lyon"


def test_secteurs_serie_et_hors_secteur():
    cs = [couple("a", 2023), couple("b", 2024), couple("c", 2024, nom_secteur=None)]
    blocs = k.secteurs(cs, [], [{"nom_secteur": "Chassieu-Bron", "ordre": 1, "wkt": SECTEUR}], 2024, 2023, NOMENC)
    assert [b["nom"] for b in blocs] == ["Chassieu-Bron", "Couples hors des secteurs"]
    assert blocs[0]["serie"] == {2023: 1, 2024: 1} and blocs[0]["lignes"][0]["statut"] == "Présence de poussins"


def test_phenologie_hivernage():
    syn = [{"annee": 2024, "premiere": date(2024, 1, 5), "derniere": date(2024, 12, 20), "nb": 3}]
    rows = k.phenologie_dates([couple("a", 2024)], [], [], syn, 2024, 2024)
    assert rows[0]["premiere_obs"] == "Hivernage" and rows[0]["derniere_obs"] == "Hivernage"
    assert rows[0]["premier_couple"] == "10/03"


def test_baguage_ignore_non_bague():
    bagues = [
        {"bague": "NONBAGUE", "classe_age": "adulte", "cle_couple": "a", "date_obs": date(2024, 4, 1)},
        {
            "bague": "E64",
            "classe_age": "poussin",
            "cle_couple": "a",
            "type_bague": "darvic",
            "date_obs": date(2024, 6, 27),
        },
    ]
    b = k.baguage(bagues, [couple("a", 2024)], 2024, NOMENC)
    assert [r["bague"] for r in b["lignes"]] == ["E64"] and b["poussins"] == 1


def test_mesures_compensatoires():
    mesures = [
        {
            "id_site": "M1",
            "id_projet": 1,
            "nom_projet": "ZAC test",
            "commune": "Saint-Priest",
            "lieu_dit": "X",
            "porteur_projet": None,
            "ref_arrete": "2011-4030",
            "surface_galet_ha": 2,
            "surface_prairie_ha": None,
            "annee_mise_en_place": 2013,
            "surface_ha": 2.1,
            "wkt": MESURE,
        }
    ]
    cs = [couple("a", 2024, id_mesure_compensatoire="M1")]
    ns = [nid("n", "a", 2024, "echec_certain", 12, id_mesure_compensatoire="M1")]
    ev = [{"id_mesure_compensatoire": "M1", "date_obs": date(2024, 4, 2), "code": "COUV", "nombre": 2}]
    m = k.mesures_compensatoires(mesures, cs, ns, ev, [], 2024, NOMENC, NOMENC["ref_statut_evenement"])
    assert m["synthese"][0]["couples"] == 1 and m["fiches"][0]["passages"][0]["evenement"] == "Couvaison observée"
    assert m["fiches"][0]["historique"][-1]["presence"] == "Oui (1)"


# ------------------------------------------------------------------------------------------------ déductions
def test_deductions_milieu_culture_lieu_dit():
    parcelle = {"wkt": "POLYGON((844000 6514000,846000 6514000,846000 6516000,844000 6516000,844000 6514000))"}
    suivi = {**parcelle, "date_suivi": date(2026, 3, 20), "type_vegetation": "maïs"}
    cs = [couple("a", 2026), couple("b", 2026, id_mesure_compensatoire="M1"), couple("c", 2026, x=830000.0)]
    ns = [nid("n", "a", 2026, "en_cours", 5)]
    stats = deduire(cs, ns, [parcelle], [suivi], [{"nom": "Les Brosses", "x": 845100.0, "y": 6515000.0}], 1000)
    assert (cs[0]["milieu"], cs[0]["culture"], cs[0]["lieu_dit"]) == ("agricole", "maïs", "Les Brosses")
    assert cs[1]["milieu"] == "parcelle_compensatoire" and cs[2]["milieu"] is None
    assert cs[0]["milieu_deduit"] and ns[0]["agricole"] and stats == {"milieu": 2, "culture": 1, "lieu_dit": 2}


# ----------------------------------------------------------------------------------------- requêtes, dossier
def test_emprise_et_schema():
    assert "ref_geo.l_areas" in PlsocQueries("s", 5).zone_sql
    assert "id_zone = 1" in PlsocQueries("s", 0, "zone_etude:1").zone_sql
    with pytest.raises(ValueError):
        PlsocQueries("s", 0, "partout")
    with pytest.raises(ValueError):
        PlsocQueries("s", 0, schema="pr_oedic_news; drop")
    assert k.annee_rapport(0, 2026) == 2026 and k.annee_rapport(2024, 2026) == 2024


def test_dossier_plsoc_pret():
    assert resolve_dossier("plsoc").name == "plsoc"
    assert check_dossier(load_dossier("plsoc")).ok


# ------------------------------------------------------------------------------------------------------ Word
def test_insertion_tableau_et_blocs(tmp_path):
    doc = Document()
    doc.add_paragraph("{{T}}")
    doc.add_paragraph("{{B}}")
    doc.add_paragraph("{{VIDE}}")
    inserer_tableau(doc, "{{T}}", [{"a": 1}, {"a": 2, "total": True}], [("a", "Colonne")])
    inserer_blocs(
        doc,
        "{{B}}",
        [
            ("titre", 3, "Secteur"),
            ("texte", "Bonjour"),
            ("a_rediger", "[À rédiger]"),
            ("tableau", [{"a": "x"}], [("a", "A")]),
        ],
    )
    inserer_tableau(doc, "{{VIDE}}", [], [("a", "A")], vide="Rien.")
    doc.save(tmp_path / "t.docx")
    d = Document(tmp_path / "t.docx")
    assert len(d.tables) == 2 and d.tables[0].rows[2].cells[0].text == "2"
    textes = [p.text for p in d.paragraphs]
    assert "Secteur" in textes and "Rien." in textes and not find_placeholders(d)


def test_rendu_complet_donnees_factices(tmp_path, monkeypatch):
    """Toutes les briques plsoc sur des données factices : aucun placeholder ne reste dans le Word."""
    fake = {
        "annee_par_defaut": 2024,
        "nomenclatures": NOMENC,
        "zone": {"surface_ha": 40000, "wkt": ZONE, "nb_communes": 3},
        "communes_carte": [{"nom": "Chassieu", "wkt": SECTEUR}],
        "secteurs": [{"nom_secteur": "Chassieu-Bron", "ordre": 1, "wkt": SECTEUR}],
        "couples": [
            couple("a", 2023),
            couple("b", 2024, id_mesure_compensatoire="M1"),
            couple("c", 2024, milieu="agricole", zonage_plu="U", statut_nidification="Nicheur probable", nb_nids=0),
        ],
        "nids": [
            nid(
                "n1",
                "b",
                2024,
                "echec_certain",
                20,
                cause_fin="predation",
                protege=True,
                type_protection="piquetage",
                id_mesure_compensatoire="M1",
            ),
            nid("n0", "a", 2023, "succes_probable", 30),
        ],
        "parcelles_mergin": [],
        "suivis_parcelle": [],
        "lieux_dits": [],
        "evenements": [{"id_mesure_compensatoire": "M1", "date_obs": date(2024, 4, 2), "code": "COUV", "nombre": 2}],
        "synthese": [{"annee": 2024, "premiere": date(2024, 2, 20), "derniere": date(2024, 11, 2), "nb": 10}],
        "rassemblements": [
            {
                "source": "mergin",
                "site": "Parcelle 1",
                "commune": "Genas",
                "annee": 2024,
                "date_comptage": date(2024, 9, 3),
                "effectif": 30,
                "date_premiere_obs": None,
                "date_derniere_obs": None,
            }
        ],
        "bagues": [
            {
                "bague": "E64",
                "classe_age": "poussin",
                "cle_couple": "b",
                "type_bague": "darvic",
                "date_obs": date(2024, 6, 27),
                "source": "mergin_poussin",
                "cle_nid": "n1",
            }
        ],
        "mesures": [
            {
                "id_site": "M1",
                "id_projet": 1,
                "nom_projet": "ZAC test",
                "commune": "Chassieu",
                "lieu_dit": "X",
                "porteur_projet": None,
                "ref_arrete": None,
                "surface_galet_ha": 0.5,
                "surface_prairie_ha": 0.5,
                "annee_mise_en_place": 2020,
                "surface_ha": 1.0,
                "wkt": MESURE,
            }
        ],
        "suivis_mesure": [],
        "anomalies": [{"gravite": "BLOQUANT", "code_anomalie": "LIEN_VIDE", "nb": 2}],
        "habitat_menace": [
            {
                "nom_epci": "Métropole de Lyon",
                "zonage_plu": "U",
                "nb_couples_annees": 1,
                "nb_parcelles": 0,
                "surface_ha": None,
                "premiere_annee": 2024,
                "derniere_annee": 2024,
            }
        ],
    }
    for nom, valeur in fake.items():
        monkeypatch.setattr(PlsocQueries, nom, lambda self, *a, v=valeur, **kw: v)
    dossier = load_dossier("plsoc")
    dirs = {n: tmp_path / n for n in ("root", "data", "tables", "maps", "dataviz", "atlas")}
    for p in dirs.values():
        p.mkdir(exist_ok=True)
    enabled, _ = dossier.select_analyses("plsoc")
    ctx = ReportContext(
        service_name="factice",
        id_area=1,
        area_name="Test",
        referee="t",
        buffer=0,
        dossier=dossier,
        enabled=enabled,
        output_dirs=dirs,
        queries=None,
        params=dict(dossier.params),
    )
    sortie = tmp_path / "rapport.docx"
    assert run_pipeline(ctx, output_file=sortie) == set()
    d = Document(sortie)
    texte = "\n".join(p.text for p in d.paragraphs)
    assert "1 couple" in texte and "Chassieu-Bron" in texte and "ZAC test" in texte
