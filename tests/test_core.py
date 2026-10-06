"""Tests du socle multi-dossiers : sans base de données ni QGIS."""

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import pytest  # noqa: E402
from docx import Document  # noqa: E402

from reportgenerator.analysis.common.models import AnalysisResult, ImageBlock, TableBlock  # noqa: E402
from reportgenerator.analysis.common.tables.generic_table import insert_general_table  # noqa: E402
from reportgenerator.core.dossier import list_dossiers, load_dossier, resolve_dossier  # noqa: E402
from reportgenerator.core.pipeline import build_plan, check_dossier  # noqa: E402
from reportgenerator.core.renderer import (apply_sections, find_placeholders,  # noqa: E402
                                           merge_results, render_report, replace_in_paragraph)


def test_replace_keeps_formatting_across_runs():
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("Parmi les ")
    bold = p.add_run("{{NB_")
    bold.bold = True
    p.add_run("TOTAL}} espèces")

    assert replace_in_paragraph(p, "{{NB_TOTAL}}", "42") == 1
    assert p.text == "Parmi les 42 espèces"
    assert p.runs[1].text == "42" and p.runs[1].bold


def test_replace_multiple_occurrences_and_value_containing_placeholder():
    doc = Document()
    p = doc.add_paragraph("{{A}} et {{A}}")
    assert replace_in_paragraph(p, "{{A}}", "{{A}}!") == 2
    assert p.text == "{{A}}! et {{A}}!"


def test_select_analyses_filters_unknown():
    dossier = load_dossier("generique")
    enabled, ignored = dossier.select_analyses("atlas_nicheur, inconnue")
    assert "atlas_nicheur" in enabled and "socle_data" in enabled
    assert ignored == ["inconnue"]


@pytest.mark.parametrize("list_analyse, expected", [
    ("analyse_eolien", "eolien"),
    ("Analyse_Eolien", "eolien"),
    ('{"analyse_eolien"}', "eolien"),
    ("atlas_nicheur", "generique"),
    ("", "generique"),
    (None, "generique"),
])
def test_resolve_dossier_from_list_analyse(list_analyse, expected):
    assert resolve_dossier(list_analyse).name == expected


def test_eolien_trigger_enables_chiropteres_and_is_not_ignored():
    dossier = resolve_dossier("analyse_eolien")
    enabled, ignored = dossier.select_analyses("analyse_eolien")
    assert ignored == [] and enabled == list(dossier.required) + ["chiropteres"]


def test_dossier_params_are_not_swallowed_by_subtables():
    params = load_dossier("eolien").params
    assert all(isinstance(v, (int, float)) for v in params["tampons_nidification_km"].values())
    assert "chiro_floutage_m" in params and "chiro_referentiel" in params


def test_cartography_runs_last():
    dossier = load_dossier("eolien")
    plan = build_plan(list(dossier.required) + ["chiropteres"])
    assert plan[-1] == "cartography" and plan.index("couches_qgis_chiro") < plan.index("cartography")


def _doc_with(*paragraphs):
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    return doc


def test_drawing_ids_unique_across_parts(tmp_path):
    """Régression : une carte recevait l'id de dessin du logo du pied de page -> Word : fichier endommagé."""
    import re
    import zipfile

    dossier = load_dossier("eolien")
    png = tmp_path / "img.png"
    plt.figure(figsize=(1, 1))
    plt.savefig(png)
    plt.close()
    keys = [k for k in find_placeholders(Document(dossier.template)) if k.startswith("Carte")]
    results = {"x": AnalysisResult(images={k: ImageBlock(png, "page") for k in keys})}
    out = tmp_path / "rapport.docx"
    render_report(dossier.template, results, out, sections=None)

    ids = []
    with zipfile.ZipFile(out) as z:
        for name in z.namelist():
            if name.startswith("word/") and name.endswith(".xml"):
                ids += re.findall(r'<wp:docPr[^>]*\bid="(\d+)"', z.read(name).decode("utf8", "ignore"))
    assert len(ids) == len(set(ids))


def test_sections_kept_or_removed():
    doc = _doc_with("avant", "{{#chiropteres}}", "chiro {{X}}", "{{/chiropteres}}", "après")
    assert apply_sections(doc, {"chiropteres"}) == []
    assert [p.text for p in doc.paragraphs] == ["avant", "chiro {{X}}", "après"]

    doc = _doc_with("avant", "{{#chiropteres}}", "chiro {{X}}", "{{/chiropteres}}", "après")
    assert apply_sections(doc, set()) == ["chiropteres"]
    assert [p.text for p in doc.paragraphs] == ["avant", "après"]
    assert find_placeholders(_doc_with("{{#a}}", "{{/a}}", "{{B}}")) == {"B"}


def test_build_plan_adds_dependencies_first():
    plan = build_plan(["synthese_generale"])
    assert plan == ["socle_data", "knowledge_status", "synthese_generale"]


def test_merge_results_rejects_duplicate_keys():
    results = {"a": AnalysisResult(texts={"X": 1}), "b": AnalysisResult(texts={"X": 2})}
    with pytest.raises(ValueError, match="deux briques"):
        merge_results(results)


def test_generique_dossier_is_complete():
    assert "generique" in list_dossiers()
    assert check_dossier(load_dossier("generique")).ok


def test_render_generique_template(tmp_path):
    dossier = load_dossier("generique")
    keys = {k for k in find_placeholders(Document(dossier.template))}

    png = tmp_path / "img.png"
    plt.figure(figsize=(1, 1))
    plt.savefig(png)
    plt.close()

    texts = {k: f"<{k}>" for k in keys if k.isupper() and not k.startswith("TABLE_")}
    tables = {
        k: TableBlock([{"a": 1}], insert_general_table, {"columns": [("A", "a")]})
        for k in keys if k.startswith("TABLE_")
    }
    images = {k.removesuffix(".png"): ImageBlock(png, "normal") for k in keys if k.endswith(".png")}
    results = {"fake": AnalysisResult(texts=texts, tables=tables, images=images)}

    output = tmp_path / "rapport.docx"
    unresolved = render_report(dossier.template, results, output)

    assert unresolved == set()
    rendered = Document(output)
    assert find_placeholders(rendered) == set()
    assert "<AREA_NAME>" in "\n".join(p.text for p in rendered.paragraphs)


def test_vm_name_unique_and_orphans():
    """Une vue par génération : deux générations simultanées ne se marchent plus dessus."""
    from datetime import datetime

    from reportgenerator.analysis.eolien.queries import VM, EolienQueries
    from reportgenerator.queries import SyntheseQueries, vm_name, vues_orphelines

    a, b = SyntheseQueries("x", 1, 5), SyntheseQueries("x", 1, 5)
    assert a.vm != b.vm and a.vm.startswith("lpoaura_afo.vm_reportgenerator_data_")

    now = datetime(2026, 10, 6, 12)
    vieille = vm_name(datetime(2026, 10, 4, 12)).split(".")[1]
    recente = vm_name(datetime(2026, 10, 6, 11)).split(".")[1]
    # l'ancien nom fixe (version précédente, éventuellement en cours ailleurs) n'est jamais supprimé
    assert vues_orphelines([vieille, recente, "vm_reportgenerator_data"], now) == [f"lpoaura_afo.{vieille}"]

    eq = EolienQueries(a)
    sql = eq._sql(f"select * from {VM} s join {VM} t using (id)")
    assert sql.count(a.vm) == 2 and eq._sql(sql) == sql
