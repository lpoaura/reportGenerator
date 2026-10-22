"""Indice de vulnérabilité des oiseaux (dossier éolien), sans base."""

import pytest

from reportgenerator.analysis.eolien import vulnerabilite as v
from reportgenerator.analysis.eolien.tables import ALL, groupes_non_vides


def _sp(nich=None, aura=None, hiv=None, migr=None, sensi=None, **kw):
    return {"lr_fr_nich": nich, "lr_aura": aura, "lr_fr_hiv": hiv, "lr_fr_migr": migr, "sensibilite": sensi, **kw}


# Lignes du tableau 2 du rapport de référence (Rapport_Eolien_VF), statuts de taxonomie.mv_c_statut
REFERENCE = [
    ("Aigle royal", _sp("VU", "VU", sensi=4), 4),
    ("Milan royal", _sp("VU", "NT", hiv="VU", migr="NA", sensi=4, hivernant=True), 4),
    ("Busard cendré", _sp("NT", "EN", migr="NA", sensi=3), 4),
    ("Goéland cendré", _sp("EN", "NA", hiv="LC", sensi=2), 3.5),       # LR France nicheur
    ("Grue cendrée", _sp("CR", None, hiv="NT", migr="NA", sensi=2), 3.5),
    ("Tadorne de Belon", _sp("LC", "CR", hiv="LC", sensi=2), 3.5),
    ("Bécassine des marais", _sp("CR", "CR", hiv="DD", migr="NA", sensi=1), 3),  # non protégée : pas de 0,5
    ("Épervier d'Europe", _sp("LC", "VU", hiv="NA", migr="NA", sensi=2), 3),
    ("Bruant ortolan", _sp("EN", "EN", migr="EN", sensi=None), 2.5),   # sensibilité inconnue = 0
    ("Faucon pèlerin", _sp("LC", "LC", hiv="NA", migr="NA", sensi=3), 2.5),
    ("Grand Cormoran", _sp("LC", "VU", hiv="LC", migr="NA", sensi=1), 2.5),
]


@pytest.mark.parametrize("nom, sp, attendu", REFERENCE, ids=[r[0] for r in REFERENCE])
def test_reproduit_le_rapport_de_reference(nom, sp, attendu):
    assert v.indice_vulnerabilite(sp) == attendu


def test_listes_migrateur_et_hivernant_seulement_si_statut_observe():
    sp = _sp("LC", "LC", hiv="EN", migr="VU", sensi=2)
    assert v.indice_vulnerabilite(sp) == 2           # ni migrateur ni hivernant sur la zone
    assert v.indice_vulnerabilite({**sp, "migrateur": True}) == 3
    assert v.indice_vulnerabilite({**sp, "hivernant": True}) == 3.5
    assert v.indice_vulnerabilite({**sp, "migrateur": True, "hivernant": True}) == 3.5


def test_notes_des_listes():
    assert v.note_liste_rouge("CR, NAm, DDw") == 5   # lr_france agrégée en repli de lr_fr_nich
    assert v.note_liste_rouge("VUm") is None         # code migrateur : lu dans lr_fr_migr
    assert v.note_liste_rouge(None) is None
    assert v.note_patrimoniale(_sp()) == v.NOTE_SANS_STATUT
    assert v.indice_vulnerabilite(_sp(sensi=0)) == 0.5
    assert v.indice_vulnerabilite(_sp("LC", lr_ra="EN", sensi=0), lr_regionale="lr_ra") == 2.5


def test_seuil_strict_ou_inclus():
    assert not v.depasse_seuil(2.5)
    assert v.depasse_seuil(3.0)
    assert v.depasse_seuil(2.5, inclus=True)
    assert not v.depasse_seuil(None)
    assert v.libelle_seuil() == "supérieur à 2,5"
    assert v.libelle_seuil(2.5, True) == "supérieur ou égal à 2,5"


def test_formats_et_couleurs():
    assert [v.fmt_note(x) for x in (3.5, 3.0, 0.5, None)] == ["3,5", "3", "0,5", ""]
    assert v.couleur(2.0) is None and v.couleur(None) is None
    assert v.couleur(2.5) == v.COULEURS[2.5] and v.couleur(4.5) == v.COULEURS[4.5]


def test_texte_tri_statut_histogramme():
    species = [
        {"nom_vern": "A", "vulnerabilite": 3.0, "code_nidif": 1, "nb_obs__all": 2},
        {"nom_vern": "B", "vulnerabilite": 4.0, "code_nidif": 0, "nb_obs__all": 1},
        {"nom_vern": "C", "vulnerabilite": 2.5, "code_nidif": 3, "nb_obs__all": 9},
    ]
    assert v.texte_vulnerabilite(species) == (
        "Au total, 2 espèces d'oiseaux obtiennent un indice de vulnérabilité supérieur à 2,5, dont 1 nicheuse "
        "(nicheur possible, probable ou certain), parmi les 3 espèces observées sur l'aire d'étude.")
    assert "3 espèces d'oiseaux" in v.texte_vulnerabilite(species, inclus=True)
    assert v.texte_vulnerabilite(species, seuil=4).startswith("Aucune des 3 espèces")
    assert [sp["nom_vern"] for sp in v.tri_vulnerabilite(species)] == ["B", "A", "C"]
    assert v.statut_biologique({"statut_nidif": "Nicheur certain", "migrateur": True}) == "Nicheur certain, migrateur"
    assert v.statut_biologique({"hivernant": True}) == "Hivernant"
    histo = {row["vulnerabilite"]: row["nb_especes"] for row in v.histogramme(species)}
    assert len(histo) == 9 and histo[3.0] == 1 and histo[0.5] == 0


def test_tableau_largeurs_imposees_et_couleurs():
    from docx import Document
    from docx.oxml.ns import qn

    from reportgenerator.analysis.common.tables.ring_table import insert_lpo_table
    doc = Document()
    doc.add_paragraph("{{T}}")
    fmt = {"v": lambda x: (v.fmt_note(x), v.couleur(x), v.couleur_texte(x))}
    insert_lpo_table(doc, "{{T}}", [{"n": "Milan royal", "v": 4.0}, {"n": "Buse", "v": 2.0}],
                     [("n", "Nom"), ("v", "Vulnérabilité")], widths=[3, 1], fixed_widths=True, cell_formats=fmt)
    table = doc.tables[0]
    assert table._tbl.tblPr.find(qn("w:tblLayout")).get(qn("w:type")) == "fixed"
    largeurs = [c._tc.tcPr.find(qn("w:tcW")).get(qn("w:w")) for c in table.rows[1].cells]
    assert largeurs == ["3750", "1250"]
    assert table.rows[1].cells[1].text == "4" and v.couleur_texte(4.0) == "FFFFFF" and v.couleur_texte(3.5) is None


def test_matrice_de_la_methode():
    lignes = {row["statut"]: row for row in v.matrice()}
    assert lignes["CR, EN = 5"]["s4"] == 4.5 and lignes["CR, EN = 5"]["s0"] == 2.5
    assert lignes["LC = 2"]["s2"] == 2 and lignes["DD, NA, non évaluée = 1"]["s0"] == 0.5


def test_phrase_espece_tres_vulnerable():
    sp = {"nom_vern": "Milan royal", "vulnerabilite": 4.0, "statut_nidif": "Nicheur certain", "migrateur": True,
          "nb_obs__all": 3293, "nb_obs__0": 2, "nb_obs__1": 10}
    d = {"dist_min_m": 0.0, "dist_nidif_m": 5812.0, "statut_nidif_proche": "Certain", "annee_nidif_proche": 2023}
    p = v.phrase_espece(sp, d, 1.0, 2016)
    assert p["nom"] == "Milan royal"
    assert p["texte"] == (" (vulnérabilité 4 ; nicheur certain, migrateur) : 3 293 données depuis 2016, dont 12 dans "
                          "la zone d'étude ou à moins de 1 km ; donnée la plus proche dans la zone d'étude ; site de "
                          "nidification certain le plus proche à 5,8 km (2023).")
    sans = v.phrase_espece({**sp, "nb_obs__0": 0, "nb_obs__1": 0, "nb_obs__all": 1, "statut_nidif": ""},
                           {"dist_min_m": 12345.0}, 1.0, 2016)
    assert "1 donnée depuis 2016, aucune dans la zone d'étude ni à moins de 1 km" in sans["texte"]
    assert "à 12,3 km de la zone d'étude" in sans["texte"]
    assert "aucune nidification probable ou certaine connue" in sans["texte"]


def test_texte_couverture_des_cartes():
    species = [{"cd_ref": 1, "nom_vern": "Vautour moine", "vulnerabilite": 4.0, "hivernant": True},
               {"cd_ref": 2, "nom_vern": "Milan royal", "vulnerabilite": 4.0},
               {"cd_ref": 3, "nom_vern": "Bruant ortolan", "vulnerabilite": 2.5}]
    texte = v.texte_couverture(species, {2})
    assert "1 espèce(s) d'indice de vulnérabilité de 3 ou plus" in texte and "Vautour moine (4 ; hivernant)" in texte
    assert "Bruant ortolan" not in texte and "Milan royal" not in texte
    assert "toutes les espèces" in v.texte_couverture(species, {1, 2})


def test_vulnerabilite_totale_par_maille():
    from reportgenerator.analysis.eolien.tables import vulnerabilite_totale
    assert vulnerabilite_totale({"vuln_especes": 3.5, "bonus_nidif": 1, "bonus_dortoir": 0,
                                 "bonus_migration": 1}) == 5.5
    assert vulnerabilite_totale({}) == 0


def test_insert_paragraph_list():
    from docx import Document

    from reportgenerator.analysis.common.paragraphs import insert_paragraph_list
    doc = Document()
    doc.add_paragraph("avant")
    doc.add_paragraph("{{LISTE}}")
    doc.add_paragraph("après")
    insert_paragraph_list(doc, "{{LISTE}}", [{"nom": "A", "texte": " : un."}, {"nom": "B", "texte": " : deux."}])
    assert [p.text for p in doc.paragraphs] == ["avant", "A : un.", "B : deux.", "après"]
    assert doc.paragraphs[1].runs[0].bold
    vide = Document()
    vide.add_paragraph("{{LISTE}}")
    insert_paragraph_list(vide, "{{LISTE}}", [])
    assert vide.paragraphs[0].text == "Aucune espèce concernée."


def test_anneaux_vides_retires():
    groups = [(ALL, "Ensemble"), (0, "Zone"), (1, "1 km"), (3, "20 km")]
    data = [{"nb_obs__all": 3, "nb_obs__3": 3}, {"nb_obs__all": 1, "nb_obs__1": 1}]
    assert groupes_non_vides(data, groups) == [(ALL, "Ensemble"), (1, "1 km"), (3, "20 km")]
    assert groupes_non_vides([], groups) == [(ALL, "Ensemble")]
