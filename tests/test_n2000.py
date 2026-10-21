"""Tests du dossier Natura 2000 (sans base)."""

import re

import pytest
from docx import Document

from reportgenerator.analysis.n2000 import selections as sel
from reportgenerator.analysis.n2000 import tables as t
from reportgenerator.bricks import n2000 as bricks
from reportgenerator.core.context import ReportContext
from reportgenerator.core.dossier import load_dossier, resolve_dossier
from reportgenerator.core.registry import get_brick
from reportgenerator.core.renderer import find_placeholders, render_report

ANNEAUX = [1.0, 10.0, 20.0]


def _pg_match(pattern, value):
    """Les motifs POSIX de selections.py, évalués comme PostgreSQL (~)."""
    return re.search(pattern, value) is not None


@pytest.mark.parametrize("n2k, oiseau, attendu", [
    ("Annexe I", True, True),
    ("Annexe I, Annexe II/2", True, True),
    ("Annexe II/2", True, False),               # annexe « chasse » de la directive Oiseaux
    ("Annexe II/1, Annexe III/2", True, False),
    ("Annexes II, IV", False, True),
    ("Annexe II", False, True),
    ("Annexe II, Annexe V", False, True),
    ("Annexe IV", False, False),
    ("Annexe V", False, False),
])
def test_definition_n2000(n2k, oiseau, attendu):
    motif = sel.N2K_ANNEXE_I_RE if oiseau else sel.N2K_ANNEXE_II_RE
    assert _pg_match(motif, n2k) is attendu


def test_rayon_par_taxon_dans_la_vm():
    from reportgenerator.bricks.common import _rayons_km
    from reportgenerator.queries import SyntheseQueries

    rayons = _rayons_km(load_dossier("natura2000").params)
    sql = SyntheseQueries._rayon_par_taxon_sql(rayons, 20000)
    assert "= 'Canis lupus' then 20000.0" in sql           # espèce et sous-espèces (binôme)
    assert "when t.ordre = 'Chiroptera' then 10000.0" in sql
    assert sql.endswith("else 1000.0 end)")
    assert SyntheseQueries._rayon_par_taxon_sql(None, 20000) == ""   # autres dossiers : pas de filtre
    with pytest.raises(ValueError):
        SyntheseQueries._rayon_par_taxon_sql({"defaut": 1, "especes": {"Gyps fulvus": 30}}, 20000)


def test_taxons_exclus():
    assert sel.exclus_sql(()) == "false"
    assert "in ('Salmo trutta')" in sel.exclus_sql(["Salmo trutta"])
    assert "and not (" in sel.n2000_sql([-60418], ["Salmo trutta"]) and "-60418" in sel.n2000_sql([-60418])


def test_complexes_n2000():
    candidats = [
        {"cd_ref": 191213, "lb_nom": "Cottus", "id_rang": "GN", "n_especes": 9, "n_especes_n2000": 9},
        {"cd_ref": 195005, "lb_nom": "Myotis", "id_rang": "GN", "n_especes": 30, "n_especes_n2000": 5},
        {"cd_ref": -60418, "lb_nom": "Myotis myotis / M. blythii", "id_rang": "RES", "n_especes": 0,
         "n_especes_n2000": 0},
        {"cd_ref": -79303, "lb_nom": "P. kuhlii / nathusii", "id_rang": "RES", "n_especes": 0, "n_especes_n2000": 0},
    ]
    retenus = t.complexes_n2000(candidats, ["Myotis myotis / M. blythii"])
    assert [(c["lb_nom"], c["origine"]) for c in retenus] == [
        ("Cottus", "automatique"), ("Myotis myotis / M. blythii", "liste du dossier")]


def test_statuts_et_presence():
    oiseau = {"groupe": "Oiseaux", "rang_nidif": 2, "migrateur": True, "hivernant": True, "nb_donnees_site": 0}
    assert t.statut_biologique(oiseau) == "Nicheur probable, migrateur, hivernant"
    assert t.statut_biologique({"groupe": "Oiseaux", "hivernant": True}) == "Hivernant"
    assert t.statut_biologique({"groupe": "Oiseaux"}) == ""   # rien ne le précise : vide
    assert t.statut_biologique({"groupe": "Amphibiens", "rang_nidif": 3}) == ""
    assert t.presence(oiseau) == "À proximité"
    assert t.presence({"nb_donnees_site": 3}) == "Dans le site"
    assert t.statut_chiro({"colonie": True, "gite": True}) == "Colonies de reproduction"
    complexe = {"nom_vern": "Murin de grande taille", "id_rang": "RES"}
    assert t.nom_affiche(complexe) == "Murin de grande taille (complexe)"


def test_textes():
    groupes = [{"groupe": "Oiseaux", "nb_donnees": 36013, "nb_especes": 180},
               {"groupe": "Amphibiens", "nb_donnees": 300, "nb_especes": 9}]
    texte = t.texte_chiffres_site(groupes, "2015-2025", [{"groupe": "Oiseaux", "nb_donnees": 50000, "nb_especes": 200}])
    assert texte.startswith("Sur la période 2015-2025, 36 313 données pour 189 espèces ont été recensées au sein "
                            "du site Natura 2000 (50 000 données pour 200 espèces en comptant le 1 kilomètre")
    assert "(180 espèces / 36 013 données)." in texte
    assert t.texte_chiffres_site([], "2015-2025").startswith("Sur la période 2015-2025, aucune donnée")
    rayons = t.texte_rayons(1, {"Accipitriformes": 10, "Falconiformes": 10, "Strigiformes": 10, "Chiroptera": 10},
                            {"Canis lupus": 20, "Lynx lynx": 20},
                            {"Canis lupus": "Loup gris", "Lynx lynx": "Lynx boréal"})
    assert "Un tampon de 1 kilomètre" in rayons
    assert "à 10 km pour les rapaces et les chiroptères" in rayons
    assert "à 20 km pour certaines espèces (Loup gris, Lynx boréal)" in rayons


def test_declencheur_et_parametres():
    dossier = load_dossier("natura2000")
    assert resolve_dossier("n2000").name == "natura2000"
    assert dossier.params["rayons_km"]["especes"]["Lynx lynx"] == 20
    assert dossier.params["groupes_taxo"] == "faune"      # pas avalé par une sous-table
    assert dossier.params["n2000_taxons_exclus"] == ["Salmo trutta"]
    rayons = [*dossier.params["rayons_km"]["ordres"].values(), *dossier.params["rayons_km"]["especes"].values()]
    assert max(rayons) <= max(dossier.params["anneaux_km"])   # emprise de la VM suffisante


TAXONS = [
    {"cd_ref": 1, "id_rang": "ES", "nom_vern": "Milan noir", "lb_nom": "Milvus migrans", "groupe": "Oiseaux",
     "ordre": "Accipitriformes", "n2k": "Annexe I", "n2000": True, "nb_donnees": 40, "nb_donnees_site": 5,
     "rang_nidif": 3, "migrateur": True, "hivernant": False, "lr_france": "LC", "lr_aura": "NT",
     "derniere_annee": 2024, "effectif_max": 12},
    {"cd_ref": 2, "id_rang": "ES", "nom_vern": "Grand Murin", "lb_nom": "Myotis myotis", "groupe": "Chauves-souris",
     "ordre": "Chiroptera", "n2k": "Annexes II, IV", "n2000": True, "nb_donnees": 15, "nb_donnees_site": 0,
     "colonie": True, "effectif_max": 111},
    {"cd_ref": 3, "id_rang": "ES", "nom_vern": "Loup gris", "lb_nom": "Canis lupus", "groupe": "Mammifères",
     "ordre": "Carnivora", "n2k": "Annexes II, IV", "n2000": True, "nb_donnees": 2, "nb_donnees_site": 0},
    {"cd_ref": 4, "id_rang": "ES", "nom_vern": "Agélène à labyrinthe", "lb_nom": "Agelena labyrinthica",
     "groupe": "Araignées", "n2000": False, "nb_donnees": 11, "nb_donnees_site": 11},
]


class FakeN2000Queries:
    def __init__(self, _):
        pass

    def site(self):
        return {"surface_ha": 877, "nb_polygones": 19}

    def complexes_candidats(self):
        return []

    def taxons(self, complexes=(), exclus=()):
        return [dict(r) for r in TAXONS]

    def chiffres_site(self, complexes=(), anneau_max=0):
        return [{"groupe": "Oiseaux", "nb_donnees": 5, "nb_especes": 1},
                {"groupe": "Araignées", "nb_donnees": 11, "nb_especes": 1}]

    def noms_vernaculaires(self, lb_noms):
        return {"Canis lupus": "Loup gris", "Lynx lynx": "Lynx boréal"}


def test_n2000_bricks_render(tmp_path, monkeypatch):
    monkeypatch.setattr(bricks, "N2000Queries", FakeN2000Queries)
    dossier = load_dossier("natura2000")

    class Q:
        anneaux_km = ANNEAUX

    ctx = ReportContext(service_name="x", id_area=1, area_name="Combe de Savoie", referee="r", buffer=1,
                        dossier=dossier, enabled=[], output_dirs={"tables": tmp_path, "data": tmp_path},
                        queries=Q(), params=dict(dossier.params))
    keys = set()
    for name in ["n2000_info", "n2000_synthese", "n2000_chiropteres", "n2000_annexes"]:
        keys |= get_brick(name).provided_keys(dossier)
        ctx.results[name] = get_brick(name).run(ctx)

    texts = {k: v for r in ctx.results.values() for k, v in r.texts.items()}
    assert texts["TITRE_RAPPORT"] == "Synthèse de données sur la zone Natura 2000 « Combe de Savoie »"
    assert texts["NB_ESPECES_N2000"] == 3 and texts["NB_OISEAUX_NICHEURS_N2000"] == 1
    assert texts["LISTE_ESPECES_GRAND_RAYON"] == "Loup gris"
    assert texts["NB_CHIRO_N2000"] == 1

    template = tmp_path / "template.docx"
    doc = Document()
    for key in sorted(keys):
        doc.add_paragraph("{{" + key + "}}")
    doc.save(template)
    assert render_report(template, ctx.results, tmp_path / "rapport.docx") == set()
    rendered = Document(tmp_path / "rapport.docx")
    assert find_placeholders(rendered) == set()
    text = "\n".join(p.text for p in rendered.paragraphs)
    # annexes : un titre par groupe présent, dans l'ordre (oiseaux, mammifères, chiroptères, invertébrés)
    titres = [p.text for p in rendered.paragraphs if p.text.startswith("Annexe ")]
    assert titres == ["Annexe 1 : liste des oiseaux recensés sur le site et à proximité",
                      "Annexe 2 : liste des mammifères recensés sur le site et à proximité",
                      "Annexe 3 : liste des chiroptères recensés sur le site et à proximité",
                      "Annexe 4 : liste des invertébrés recensés sur le site et à proximité"]
    assert "877" in text
    assert (tmp_path / "1_especes_n2000.xlsx").exists() and (tmp_path / "annexe_4_invertébrés.xlsx").exists()


def test_sources_et_cles():
    sources = [{"source": "faune-france", "n": 900}, {"source": "gn2pg_pnr_bauges", "n": 50},
               {"source": "gn2pg_cen38", "n": 40}, {"source": "MONITORING_STOM", "n": 5}]
    assert t.texte_sources(sources) == "VisioNature (Réseau LPO), PNR Bauges, CEN 38 et autres partenaires"
    assert t.libelle_source("Import", {"Import": "Partenaires (imports)"}) == "Partenaires (imports)"
    assert t.cle("Martin-pêcheur d'Europe") == "martin_pecheur_d_europe"
    assert t.ordre_groupe("Oiseaux") < t.ordre_groupe("Odonates") < t.ordre_groupe("Araignées")


def test_taxons_atlas():
    taxons = [{"n2000": True, "groupe": "Oiseaux", "ordre": "Accipitriformes", "nom": "Milan noir", "rang_nidif": 3},
              {"n2000": True, "groupe": "Oiseaux", "ordre": "Ciconiiformes", "nom": "Grande Aigrette"},
              {"n2000": True, "groupe": "Chauves-souris", "ordre": "Chiroptera", "nom": "Grand Murin"},
              {"n2000": True, "groupe": "Odonates", "ordre": "Odonata", "nom": "Agrion de Mercure"},
              {"n2000": False, "groupe": "Odonates", "ordre": "Odonata", "nom": "Agrion jouvencelle"}]
    assert [r["nom"] for r in t.taxons_atlas(taxons)] == ["Milan noir", "Agrion de Mercure"]
    assert "Grande Aigrette" in [r["nom"] for r in t.taxons_atlas(taxons, oiseaux_non_nicheurs=True)]


def test_liste_de_figures(tmp_path):
    """Atlas dans le Word : titre de groupe, carte, légende numérotée, texte à rédiger ; figure absente ignorée."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    from reportgenerator.analysis.common.figures import insert_figure_list

    png = tmp_path / "carte.png"
    plt.figure(figsize=(2, 3)).savefig(png)
    doc = Document()
    doc.add_paragraph("{{ATLAS}}")
    figures = [{"path": png, "legende": "Répartition – Milan noir"},
               {"path": tmp_path / "absente.png", "legende": "x"}]
    data = [{"titre": "Oiseaux nicheurs", "figures": figures},
            {"titre": "Reptiles", "figures": [{"path": tmp_path / "absente.png", "legende": "y"}]}]
    insert_figure_list(doc, "{{ATLAS}}", data)
    textes = [p.text for p in doc.paragraphs]
    assert textes == ["Oiseaux nicheurs", "", "Figure  : Répartition – Milan noir",
                      "[Photo et texte de l'espèce à rédiger]"]
    assert len(doc.inline_shapes) == 1 and doc.paragraphs[0].style.name == "Heading 3"
    assert 'SEQ Figure' in doc.paragraphs[2]._p.xml
    # numéro affiché sans mise à jour des champs : figures déjà présentes + 1 ; avertissement en italique
    doc.add_paragraph("{{ATLAS2}}")
    insert_figure_list(doc, "{{ATLAS2}}", [{"titre": "Odonates", "avertissement": "À lire avec prudence",
                                            "figures": [{"path": png, "legende": "z"}]}], note=None)
    seq = [p for p in doc.paragraphs if "SEQ Figure" in p._p.xml]
    assert [p._p.xpath(".//w:fldSimple//w:t")[0].text for p in seq] == ["1", "2"]
    alerte = next(p for p in doc.paragraphs if p.text == "À lire avec prudence")
    assert alerte.runs[0].italic and alerte.paragraph_format.keep_with_next


def test_groupes_pour_graphique():
    groupes = [{"groupe": f"G{i}", "nb_donnees": 100 - i, "nb_especes": 10} for i in range(20)]
    rows = t.groupes_pour_graphique(groupes, max_groupes=15)
    assert len(rows) == 15 and rows[0] == {"group_taxo": "G0", "nb_espece": 10, "nb_data_tot": 100}
    assert rows[-1] == {"group_taxo": "Autres groupes (6)", "nb_espece": 60,
                        "nb_data_tot": sum(100 - i for i in range(14, 20))}
    # un seul groupe en trop : pas de regroupement
    assert len(t.groupes_pour_graphique(groupes[:15], max_groupes=15)) == 15


# (état des connaissances par maille : tests/test_connaissance.py)
