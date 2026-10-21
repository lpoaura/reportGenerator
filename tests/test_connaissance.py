"""État des connaissances par maille (module et briques communs aux dossiers), sans base ni QGIS."""

from docx import Document

from reportgenerator.analysis.connaissance import indicateurs as ind
from reportgenerator.analysis.connaissance.selections import enjeu_sql
from reportgenerator.analysis.n2000 import tables as t
from reportgenerator.bricks import connaissance as bricks
from reportgenerator.core.context import ReportContext
from reportgenerator.core.dossier import load_dossier
from reportgenerator.core.registry import get_brick
from reportgenerator.core.renderer import find_placeholders, render_report

MAILLES = [
    {"groupe": "Odonates", "id_area": 1, "nb_donnees": 1, "nb_jours": 1, "nb_especes": 1, "nb_enjeux": 0},
    {"groupe": "Odonates", "id_area": 2, "nb_donnees": 40, "nb_jours": 12, "nb_especes": 30, "nb_enjeux": 2},
    {"groupe": "Reptiles", "id_area": 2, "nb_donnees": 5, "nb_jours": 4, "nb_especes": 3, "nb_enjeux": 0},
]


def test_indicateurs_connaissance():
    m = ind.indicateurs_mailles(MAILLES, {"Odonates": 57, "Reptiles": 12})
    assert m[0]["peu_donnees"] and not m[1]["peu_donnees"]
    assert m[1]["pct_pool_local"] == 52.6 and m[1]["classe_croisee"] == "e2r2"      # riche et bien prospecté
    assert m[0]["classe_croisee"] == "e0r0" and m[2]["classe_croisee"] == "e1r1"    # 25 % du pool : moyen
    # carte dès 4 espèces dans la zone, quel que soit l'effort ; l'effort ponctuel donne un avertissement
    assert ind.groupes_cartographies(m, {"Odonates": 4, "Reptiles": 3}) == ["Odonates"]
    # (dans l'ordre des groupes du rapport)
    assert ind.groupes_cartographies(m, {"Odonates": 4, "Reptiles": 3}, min_especes_site=3) == ["Reptiles", "Odonates"]
    assert ind.medianes_jours(m) == {"Odonates": 6.5, "Reptiles": 4}
    assert ind.medianes_jours(m[:1] + m[2:]) == {"Odonates": 1, "Reptiles": 4}
    assert ind.texte_avertissement_effort(6.5) is None
    assert "médiane : 1,5 jour d'observation" in ind.texte_avertissement_effort(1.5)
    rows = ind.synthese_groupes(m, {"Odonates": 50, "Reptiles": 9}, {"Odonates": 57, "Reptiles": 12},
                                {"Odonates": 84, "Reptiles": 42}, ["Odonates"], especes_site={"Odonates": 30})
    odo = next(r for r in rows if r["groupe"] == "Odonates")
    assert (odo["nb_especes_site"], odo["pct_local"], odo["pct_regional"], odo["mailles_peu_donnees"],
            odo["jours_median"], odo["carte"]) == (30, "88 %", "60 %", 1, "6,5", "oui")
    # seuil « peu de données » paramétrable
    rows6 = ind.synthese_groupes(m, {"Odonates": 50}, {}, {}, [], seuil_donnees=6)
    assert rows6[0]["mailles_peu_donnees"] == 1 and rows6[0]["carte"] == "non"


def test_textes_connaissance():
    txt = ind.texte_groupes_cartographies(["Oiseaux", "Papillons de jour", "Odonates"], 4, {"Odonates": 1}, 2)
    assert "au moins 4 espèces connues dans le site" in txt and "oiseaux, papillons de jour et odonates" in txt
    assert "Pour odonates, la moitié" in txt and "2 jours" in txt
    assert "avertissement" not in ind.texte_groupes_cartographies(["Oiseaux"], 4, {}, 2)
    assert ind.texte_groupes_cartographies([], 4).startswith("Aucun groupe")
    # vocabulaire du dossier (générique, éolien)
    gen = ind.texte_groupes_cartographies(["Oiseaux"], 4, zone="la zone d'étude", renvoi="les tableaux du rapport")
    assert "dans la zone d'étude" in gen and "dans les tableaux du rapport" in gen
    assert "1 à 2 jours, 3 à 9 jours, 10 jours et plus" in ind.texte_lecture_croisee()
    assert "moins de 25 %, 25 à 50 %, plus de 50 %" in ind.texte_lecture_croisee()
    assert ind.texte_emprise("le site", 1) == "le site et le kilomètre autour"
    assert ind.texte_emprise("la zone d'étude", 2.5) == "la zone d'étude et les 2,5 km autour"
    assert ind.texte_emprise("la zone d'étude", 0) == "la zone d'étude"


def test_emprises():
    """Cadrage des cartes : contour du rayon du taxon (règle de la vue) + marge."""
    from reportgenerator.bricks.common import _rayons_km
    rayons = _rayons_km(dict(load_dossier("natura2000").params))
    assert t.rayon_taxon_km(rayons, "Pernis apivorus", "Accipitriformes") == 10       # rapace
    assert t.rayon_taxon_km(rayons, "Canis lupus italicus", "Carnivora") == 20         # sous-espèce du loup
    assert t.rayon_taxon_km(rayons, "Myotis myotis / M. blythii", "Chiroptera") == 10  # complexe
    assert t.rayon_taxon_km(rayons, "Bombina variegata", "Anura") == 1                 # défaut
    assert ind.marge_m(1000) == 500 and ind.marge_m(10000) == 1000
    bbox = [900000, 6500000, 910000, 6505000]
    assert ind.emprise_contour(bbox, 1500) == [898500, 6498500, 911500, 6506500]
    assert ind.rectangle_wkt([0, 0, 2, 1]) == "POLYGON((0 0,2 0,2 1,0 1,0 0))"


def test_enjeu_sql():
    assert "Possible" in enjeu_sql(True) and "Possible" not in enjeu_sql(False)
    assert "('CR', 'EN', 'VU', 'NT')" in enjeu_sql()


class FakeQueries:
    """Requêtes de l'état des connaissances (sans base) : 2 groupes, dont 1 cartographié avec avertissement."""
    appels = {}

    def __init__(self, _):
        pass

    def pool_local(self, rayon_km, annees=None, statuts_validation=()):
        FakeQueries.appels["pool"] = (rayon_km, annees)
        return {"Oiseaux": 120, "Chauves-souris": 20}

    def mailles(self, rayon_m, code_grille, groupes=(), enjeu_sql="false"):
        FakeQueries.appels["mailles"] = (rayon_m, code_grille, groupes)
        return [{"groupe": "Oiseaux", "id_area": i, "nb_donnees": 3, "nb_jours": 1, "nb_especes": 10,
                 "nb_enjeux": 1, "geom": "POLYGON((0 0,1 0,1 1,0 1,0 0))"} for i in range(5)]

    def especes_par_groupe(self, rayon_m, groupes=()):
        return [{"groupe": "Oiseaux", "nb_especes_zone": 30, "nb_especes": 60},
                {"groupe": "Chauves-souris", "nb_especes_zone": 2, "nb_especes": 5}]

    def grille(self, code_grille, rayon_m):
        return []

    def enveloppe_zone(self):
        return [900000, 6500000, 902000, 6501000]


def _ctx(tmp_path, nom, enabled=(), buffer=2):
    dossier = load_dossier(nom)

    class Q:
        anneaux_km = [float(a) for a in dossier.params.get("anneaux_km", [buffer])]

        def get_number_esp_per_taxonomy(self):
            return [{"tx_group2_inpn_v2": "Oiseaux", "nb_esp": 450}]

    for d in ("tables", "data", "maps"):
        (tmp_path / d).mkdir(exist_ok=True)
    return ReportContext(service_name="x", id_area=1, area_name="Test", referee="r", buffer=buffer,
                         dossier=dossier, enabled=list(enabled), output_dirs={d: tmp_path / d for d in
                                                                              ("tables", "data", "maps")},
                         queries=Q(), params=dict(dossier.params))


def test_briques_connaissance_generique(tmp_path, monkeypatch):
    monkeypatch.setattr(bricks, "ConnaissanceQueries", FakeQueries)
    ctx = _ctx(tmp_path, "generique", buffer=2)
    keys = set()
    for name in ["connaissance", "couches_qgis_connaissance"]:
        keys |= get_brick(name).provided_keys(ctx.dossier)
        ctx.results[name] = get_brick(name).run(ctx)
    # générique : mailles sur zone + buffer du formulaire, pool à 10 km, toutes les années, groupes de la VM
    assert FakeQueries.appels["mailles"] == (2000.0, "M1", ())
    assert FakeQueries.appels["pool"] == (10.0, None)
    res = ctx.results["connaissance"]
    assert res.data["groupes"] == ["Oiseaux"]                       # chiroptères : 2 espèces dans la zone
    assert "Oiseaux" in res.data["avertissements"]                  # 1 jour par maille
    assert res.texts["CONNAISSANCE_EMPRISE"] == "la zone d'étude et les 2 km autour"
    assert "dans la zone d'étude" in res.texts["TEXTE_GROUPES_CARTOGRAPHIES"]
    render = ctx.params["qgis_render"]
    assert render["layout_extents"][bricks.EFFORT_DONNEES] == [897000.0, 6497000.0, 905000.0, 6504000.0]
    assert render["project_variables"]["pool_local_km"] == "10"
    assert (tmp_path / "data" / "connaissance_atlas_groupes.gpkg").exists()
    assert (tmp_path / "tables" / "connaissance_groupes.xlsx").exists()

    # rendu Word des sorties du module
    template = tmp_path / "template.docx"
    doc = Document()
    for key in sorted(keys):
        doc.add_paragraph("{{" + key + "}}")
    doc.save(template)
    assert render_report(template, ctx.results, tmp_path / "rapport.docx") == set()
    assert find_placeholders(Document(tmp_path / "rapport.docx")) == set()


def test_briques_connaissance_eolien(tmp_path, monkeypatch):
    """Éolien : oiseaux, et chiroptères seulement si l'analyse optionnelle est activée ; emprise 6 km."""
    monkeypatch.setattr(bricks, "ConnaissanceQueries", FakeQueries)
    ctx = _ctx(tmp_path, "eolien")
    get_brick("connaissance").run(ctx)
    rayon_m, _, groupes = FakeQueries.appels["mailles"]
    assert groupes == ("Oiseaux",) and rayon_m == ctx.params["connaissance_rayon_km"] * 1000
    assert FakeQueries.appels["pool"][0] >= ctx.params["connaissance_rayon_km"]
    ctx = _ctx(tmp_path, "eolien", enabled=["chiropteres"])
    get_brick("connaissance").run(ctx)
    assert FakeQueries.appels["mailles"][2] == ("Oiseaux", "Chauves-souris")
