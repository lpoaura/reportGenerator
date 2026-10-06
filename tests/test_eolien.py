"""Tests du dossier éolien sans base ni QGIS (requêtes remplacées par des données factices)."""

from datetime import datetime

from docx import Document

from reportgenerator.analysis.cartography.qgis_sources import relink_gpkg_source
from reportgenerator.analysis.eolien.tables import (anneaux_texte, est_enjeu, periode_etude, phrase_precautions,
                                                    pivot_par_anneau, resume_par_anneau,
                                                    sensibilite_totale, texte_connaissance)
from reportgenerator.bricks import eolien as bricks
from reportgenerator.core.context import ReportContext
from reportgenerator.core.dossier import load_dossier
from reportgenerator.core.pipeline import build_plan
from reportgenerator.core.registry import get_brick
from reportgenerator.core.renderer import find_placeholders, render_report

ANNEAUX = [1.0, 6.0, 20.0]


def _row(cd_ref, anneau, nb, annee, **kw):
    base = {"cd_ref": cd_ref, "anneau_ordre": anneau, "nb_obs": nb, "derniere_annee": annee,
            "effectif_max": kw.pop("effectif", 1), "code_nidif": kw.pop("nidif", 0),
            "nom_vern": f"Espèce {cd_ref}", "lb_nom": f"Genus sp{cd_ref}", "famille": "Fringillidae",
            "ordre": "Passeriformes", "lr_aura": "LC", "lr_fr_nich": "LC", "sensibilite": 0}
    base.update(kw)
    return base


ROWS = [
    _row(1, 0, 3, 2020),
    _row(1, 2, 5, 2024, nidif=2),
    _row(2, 3, 10, 2023, ordre="Accipitriformes", lr_aura="VU", sensibilite=3, nidif=3, effectif=4),
    _row(3, 1, 1, 2019, lr_fr_nich="LC, VUm"),
]


def test_pivot_sums_disjoint_rings():
    species = pivot_par_anneau(ROWS)
    sp1 = next(sp for sp in species if sp["cd_ref"] == 1)
    assert sp1["nb_obs__all"] == 8
    assert sp1["derniere_annee__all"] == 2024
    assert sp1["statut_nidif"] == "Nicheur probable"
    assert "nb_obs__1" not in sp1


def test_est_enjeu_applies_regional_red_list_and_ignores_migrant_codes():
    species = {sp["cd_ref"]: sp for sp in pivot_par_anneau(ROWS)}
    assert est_enjeu(species[2])
    assert not est_enjeu(species[1])
    # "VUm" (migrateur) ne rend pas l'espèce "à enjeux nicheur"
    assert not est_enjeu(species[3])
    assert est_enjeu({**species[1], "lr_aura": "EN"})


def test_resume_and_text():
    species = pivot_par_anneau(ROWS)
    resume = resume_par_anneau(species, ANNEAUX)
    by_ring = {r["anneau"]: r for r in resume}
    assert by_ring[0] == {"anneau": 0, "zone": "Zone d'étude", "nb_obs": 3, "nb_especes": 1, "derniere_annee": 2020}
    assert by_ring["all"]["nb_obs"] == 19 and by_ring["all"]["nb_especes"] == 3
    text = texte_connaissance(resume, ANNEAUX)
    assert "rayon de 20 kilomètres" in text and "Entre 6 et 20 km" in text


def test_small_formatters():
    assert periode_etude(10, datetime(2026, 5, 1)) == "(2016 à 2026)"
    assert anneaux_texte([1.0, 6.0, 20.0]) == "1 km, 6 km et 20 km"
    assert sensibilite_totale({"sensi_especes": 4, "bonus_nidif": 1, "bonus_dortoir": 0, "bonus_migration": 1}) == 6


def test_phrases_anneaux_accords():
    from reportgenerator.analysis.eolien.tables import phrases_anneaux

    par_anneau = {3: {"nb_obs": 2, "nb_especes": 0}, 2: {"nb_obs": 1, "nb_especes": 1},
                  1: {"nb_obs": 0, "nb_especes": 0}, 0: {"nb_obs": 3, "nb_especes": 1}}
    phrases = phrases_anneaux(par_anneau, ANNEAUX, "contactées", "observations", unite_anneaux="observation")
    assert phrases == [
        "Entre 6 et 20 km autour du projet, on compte 2 observations, sans identification à l'espèce.",
        "Entre 1 et 6 km autour du projet, on compte 1 observation concernant 1 espèce.",
        "Entre 0 et 1 km autour du projet, aucune donnée n'est connue.",
        "Dans la zone d'étude immédiate du projet, 1 espèce a été contactée pour 3 observations.",
    ]


def test_phrase_precautions():
    texte = phrase_precautions(37.1, "Auvergne-Rhône-Alpes", 100, 22)
    assert "37 % de l'aire d'étude se situe hors de la région Auvergne-Rhône-Alpes" in texte
    assert "22 % des données sont localisées au lieu-dit" in texte
    assert phrase_precautions(0.2, "AuRA", 100, 1) == ""
    assert phrase_precautions(None, "AuRA", 0, 0) == ""


def test_relink_keeps_subset():
    src = './data/donnees_brutes.gpkg|layername=donnees_brutes|subset="mortality" = \'True\''
    assert relink_gpkg_source(src, "/out/data") == (
        '/out/data/donnees_brutes.gpkg|layername=donnees_brutes|subset="mortality" = \'True\''
    )
    assert relink_gpkg_source("service='x' table=\"a\".\"b\"", "/out") is None


def test_tampons_params_have_homogeneous_types(monkeypatch):
    """Régression : psycopg refuse une liste mêlant int et float (5 et 2.5 du dossier.toml)."""
    from reportgenerator.analysis.eolien.queries import EolienQueries

    captured = {}
    queries = EolienQueries(FakeSyntheseForQueries())
    monkeypatch.setattr(queries, "_fetch", lambda sql, params=None: captured.setdefault("params", params) and [])

    tampons = load_dossier("eolien").params["tampons_nidification_km"]
    queries.nidification_tampons(tampons)

    noms, rayons, statuts = captured["params"]
    assert all(isinstance(n, str) for n in noms)
    assert all(isinstance(r, float) for r in rayons)
    assert rayons[noms.index("Milvus milvus")] == 5.0


class FakeSyntheseForQueries:
    service_name = "x"
    id_area = 1
    anneaux_km = ANNEAUX


class FakeEolienQueries:
    def __init__(self, _):
        pass

    def species_ring_stats(self, where="true"):
        return ROWS

    def points(self, where, extra_columns=""):
        return [{"id_synthese": 1, "geom": "POINT(700000 6500000)"}]

    def perimetres(self):
        return []

    def nidification_tampons(self, tampons_km, statuts=("Certain",)):
        return []

    def sensibilite_mailles(self):
        return [{"geom_maille": None, "id_area": 9, "sensi_especes": 3,
                 "bonus_nidif": 1, "bonus_dortoir": 0, "bonus_migration": 0}]

    def fetch(self, sql):
        return [{"id": 1, "geom": "POINT(700000 6500000)"}]

    def hors_region(self, region):
        return {"pct_hors_region": 37.1, "geom": "POLYGON((0 0, 1 0, 1 1, 0 0))"}

    def qualite_donnees(self):
        return {"oiseaux": {"nb_donnees": 100, "nb_lieu_dit": 22, "nb_cachees": 16, "nb_observations": 90},
                "chiro": {"nb_donnees": 500, "nb_lieu_dit": 0, "nb_cachees": 0, "nb_observations": 6}}


class FakeSynthese:
    anneaux_km = ANNEAUX

    def get_knowledge_protected_area(self, rayon_km=None):
        return []


def test_eolien_bricks_render(tmp_path, monkeypatch):
    monkeypatch.setattr(bricks, "EolienQueries", FakeEolienQueries)
    layers = []
    monkeypatch.setattr(bricks, "_export_layer", lambda ctx, rows, name, geom_col="geom": layers.append(name))

    dossier = load_dossier("eolien")
    dirs = {"tables": tmp_path, "data": tmp_path}
    ctx = ReportContext(
        service_name="x", id_area=1, area_name="Test (07)", referee="r", buffer=5, dossier=dossier,
        enabled=[], output_dirs=dirs, queries=FakeSynthese(),
        params=dict(dossier.params),
    )
    names = [n for n in build_plan(list(dossier.required)) if n not in ("socle_data", "cartography")]
    keys = set()
    for name in names:
        keys |= get_brick(name).provided_keys(dossier)
        ctx.results[name] = get_brick(name).run(ctx)

    template = tmp_path / "template.docx"
    doc = Document()
    for key in sorted(keys):
        doc.add_paragraph("{{" + key + "}}")
    doc.save(template)

    unresolved = render_report(template, ctx.results, tmp_path / "rapport.docx")
    assert unresolved == set()

    rendered = Document(tmp_path / "rapport.docx")
    assert find_placeholders(rendered) == set()
    text = "\n".join(p.text for p in rendered.paragraphs)
    assert "Test (07)" in text and "(20" in text
    # tableau espèces à enjeux : 2 lignes d'en-tête + 1 espèce (cd_ref 2)
    enjeux_table = rendered.tables[[i for i, k in enumerate(sorted(k for k in keys if k.startswith("TABLE_")))
                                    if k == "TABLE_ESPECES_ENJEUX"][0]]
    assert len(enjeux_table.rows) == 3
    # couches attendues par le projet QGIS modèle (noms de l'ancien workflow R)
    assert set(layers) == {
        "vm_eolienne_zone_etude", "vm_eolienne_limit", "vm_eolienne_zone_protection", "vm_eolienne",
        "vm_eolienne_etat_connaissance", "vm_eolienne_nidif", "vm_eolienne_nidif_cn",
        "vm_eolienne_nidif_tampon", "vm_eolienne_dortoirs", "vm_eolienne_migration_all",
        "vm_eolienne_sensi_oiseau", "vm_eolienne_hors_region",
    }
    assert "Précautions de lecture" in ctx.results["perimetres_anneaux"].texts["CONNAISSANCE_OISEAUX"]
    sensi = ctx.results["couches_qgis_eolien"].data
    assert sensi["vm_eolienne_sensi_oiseau"] == 1


def test_compat_sql_keeps_r_field_names():
    from reportgenerator.analysis.eolien import couches_qgis as c

    assert "nb_annee" in c.NIDIF_SQL and "s.vn_nom_fr as nom_vern" in c.NIDIF_SQL
    assert "esp_enjeux" in c.MIGRATION_SQL and "nb_ind_max" in c.MIGRATION_SQL
    assert "sum_nb_data" in c.CONNAISSANCE_SQL and "sum_nb_esp" in c.CONNAISSANCE_SQL
    limit = c.limit_sql(1, ANNEAUX)
    assert "'contour 1 km'" in limit and "'contour 20 km'" in limit and "20000.0" in limit
    assert "name_project" in c.zone_etude_sql(1, "L'Étang")
    assert "'L''Étang'" in c.zone_etude_sql(1, "L'Étang")
