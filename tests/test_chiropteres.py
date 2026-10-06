"""Tests de la partie chiroptères du dossier éolien (sans base ni QGIS)."""

import pytest
from docx import Document

from reportgenerator.analysis.eolien.chiropteres import (classe, liste_vulnerables, load_referentiel,
                                                         note_patrimoniale, pivot_gites, texte_gites,
                                                         vulnerabilite)
from reportgenerator.analysis.eolien.tables import ALL, pivot_par_anneau, resume_par_anneau
from reportgenerator.bricks import chiropteres as chiro_bricks
from reportgenerator.bricks import eolien as eolien_bricks
from reportgenerator.core.context import ReportContext
from reportgenerator.core.dossier import load_dossier
from reportgenerator.core.registry import get_brick
from reportgenerator.core.renderer import find_placeholders, render_report

DOSSIER = load_dossier("eolien")
REF = load_referentiel(DOSSIER.path / DOSSIER.params["chiro_referentiel"])
SEUILS_NE = DOSSIER.params["chiro_seuils_ne"]
SEUILS_NH = DOSSIER.params["chiro_seuils_nh"]

# Tableau 5 du document chiroptères : (cd_ref, lr_aura, n2k, NH, vuln. habitat, NE, vuln. éolien)
TABLEAU_5 = [
    (79305, "EN", "Annexes II, IV", 8, "Forte", 10, "Très forte"),      # Minioptère de Schreibers
    (79301, "VU", "Annexes II, IV", 9, "Forte", 6, "Moyenne"),          # Murin de Bechstein
    (60427, "VU", "Annexes II, IV", 6, "Moyenne", 9, "Forte"),          # Petit Murin
    (60490, "DD", "Annexe IV", 6, "Moyenne", 8, "Forte"),               # Pipistrelle de Nathusius
    (60345, "LC", "Annexes II, IV", 8, "Forte", 5, "Moyenne"),          # Barbastelle d'Europe
    (60418, "LC", "Annexes II, IV", 5, "Moyenne", 8, "Forte"),          # Grand Murin
    (60400, "NT", "Annexes II, IV", 7, "Moyenne", 5, "Moyenne"),        # Murin à oreilles échancrées
    (60468, "VU", "Annexe IV", 5, "Moyenne", 8, "Forte"),               # Noctule commune
    (60479, "NT", "Annexe IV", 6, "Moyenne", 7, "Forte"),               # Pipistrelle commune
    (79303, "LC", "Annexe IV", 6, "Moyenne", 7, "Forte"),               # Pipistrelle de Kuhl
    (60489, "LC", "Annexe IV", 6, "Moyenne", 7, "Forte"),               # Pipistrelle pygmée
    (60295, "NT", "Annexes II, IV", 9, "Forte", 3, "Faible"),           # Grand rhinolophe
    (60461, "LC", "Annexe IV", 4, "Faible", 7, "Forte"),                # Noctule de Leisler
    (60313, "LC", "Annexes II, IV", 8, "Forte", 3, "Faible"),           # Petit rhinolophe
    (60506, "LC", "Annexe IV", 6, "Moyenne", 6, "Moyenne"),             # Vespère de Savi
    (79299, "DD", "Annexe IV", 8, "Forte", 2, "Faible"),                # Murin d'Alcathoé
    (60408, "DD", "Annexe IV", 8, "Forte", 2, "Faible"),                # Murin de Natterer
    (60537, "DD", "Annexe IV", 2, "Faible", 8, "Forte"),                # Sérotine bicolore
    (60383, "LC", "Annexe IV", 7, "Moyenne", 2, "Faible"),              # Murin à moustaches
    (60360, "NT", "Annexe IV", 2, "Faible", 7, "Forte"),                # Sérotine commune
    (60527, "LC", "Annexe IV", 6, "Moyenne", 4, "Faible"),              # Oreillard gris
    (60518, "LC", "Annexe IV", 6, "Moyenne", 4, "Faible"),              # Oreillard roux
    (60557, "NT", "Annexe IV", 2, "Faible", 6, "Moyenne"),              # Molosse de Cestoni
    (200118, "LC", "Annexe IV", 4, "Faible", 4, "Faible"),              # Murin de Daubenton
]


@pytest.mark.parametrize("cd_ref, lr_aura, n2k, nh, nh_classe, ne, ne_classe", TABLEAU_5)
def test_vulnerabilite_reproduit_le_tableau_du_document(cd_ref, lr_aura, n2k, nh, nh_classe, ne, ne_classe):
    sp = vulnerabilite({"cd_ref": cd_ref, "lr_aura": lr_aura, "n2k": n2k}, REF, SEUILS_NE, SEUILS_NH)
    assert (sp["nh"], sp["nh_classe"], sp["ne"], sp["ne_classe"]) == (nh, nh_classe, ne, ne_classe)


def test_grande_noctule_ecart_documente():
    """Le document donne NE = 8 ; la méthode (NSE 8, NP 4) donne 7. La classe reste « Forte »."""
    sp = vulnerabilite({"cd_ref": 60457, "lr_aura": "VU", "n2k": "Annexe IV"}, REF, SEUILS_NE, SEUILS_NH)
    assert (sp["ne"], sp["ne_classe"], sp["nh"]) == (7, "Forte", 5)


def test_note_patrimoniale_et_classes():
    assert note_patrimoniale("EN", "Annexes II, IV") == 9
    assert note_patrimoniale("LC", "Annexe IV") == 1
    assert note_patrimoniale(None, None) == 1
    assert classe(10, SEUILS_NE) == "Très forte" and classe(4, SEUILS_NE) == "Faible"
    assert classe(None, SEUILS_NE) == ""


def test_complexes_comptes_en_donnees_pas_en_especes():
    rows = [
        {"cd_ref": 60479, "anneau_ordre": 1, "nb_obs": 5, "derniere_annee": 2024, "id_rang": "ES"},
        {"cd_ref": 196414, "anneau_ordre": 1, "nb_obs": 3, "derniere_annee": 2023, "id_rang": "GN"},
    ]
    resume = {r["anneau"]: r for r in resume_par_anneau(pivot_par_anneau(rows), [1, 6, 20])}
    assert resume[ALL]["nb_obs"] == 8 and resume[ALL]["nb_especes"] == 1


def test_gites_et_texte():
    rows = [
        {"cd_ref": 60313, "anneau_ordre": 2, "type_gite": "Colonie de reproduction", "nom_vern": "Petit rhinolophe",
         "lb_nom": "Rhinolophus hipposideros", "id_rang": "ES", "nb_sites": 3, "effectif_max": 40,
         "derniere_annee": 2024},
        {"cd_ref": 60313, "anneau_ordre": 3, "type_gite": "Colonie de reproduction", "nom_vern": "Petit rhinolophe",
         "lb_nom": "Rhinolophus hipposideros", "id_rang": "ES", "nb_sites": 2, "effectif_max": 12,
         "derniere_annee": 2023},
        {"cd_ref": 60518, "anneau_ordre": 3, "type_gite": "Gîte", "nom_vern": "Oreillard roux",
         "lb_nom": "Plecotus auritus", "id_rang": "ES", "nb_sites": 1, "effectif_max": 2, "derniere_annee": 2020},
    ]
    gites = pivot_gites(rows)
    assert gites[0]["type_gite"] == "Colonie de reproduction"
    assert gites[0][f"nb_sites__{ALL}"] == 5 and gites[0]["effectif_max"] == 40
    texte = texte_gites({"nb_colonies": 5, "nb_gites": 1, "annee_min_colonies": 2019, "annee_max_colonies": 2024},
                        gites, 20)
    assert "Petit rhinolophe (5)" in texte and "1 site de gîte hors reproduction est connu" in texte
    assert "Aucune colonie" in texte_gites({"nb_colonies": 0, "nb_gites": 0}, [], 20)


@pytest.mark.parametrize("enabled, titre", [
    (["chiropteres"], "Synthèse des données d’oiseaux et de chiroptères dans le cadre d’un projet éolien"),
    ([], "Synthèse des données d’oiseaux dans le cadre d’un projet éolien"),
])
def test_titre_selon_analyses(enabled, titre):
    ctx = ReportContext(service_name="x", id_area=1, area_name="Z", referee="r", buffer=5, dossier=DOSSIER,
                        enabled=enabled, output_dirs={}, queries=None, params=dict(DOSSIER.params))
    texts = get_brick("projet_info").run(ctx).texts
    assert texts["TITRE_RAPPORT"] == titre
    assert ("chiroptérologiques" in texts["ENJEUX_ETUDIES"]) == bool(enabled)


def test_liste_vulnerables():
    donnees = [(79305, "EN", "Annexes II, IV", "Minioptère"), (60313, "LC", "Annexes II, IV", "Petit rhinolophe")]
    especes = [vulnerabilite({"cd_ref": c, "lr_aura": lr, "n2k": n, "nom_vern": nom}, REF, SEUILS_NE, SEUILS_NH)
               for c, lr, n, nom in donnees]
    assert liste_vulnerables(especes, "ne_classe") == "Minioptère (très forte)"
    assert liste_vulnerables(especes, "nh_classe") == "Minioptère (forte), Petit rhinolophe (forte)"


class FakeChiroQueries:
    def __init__(self, _):
        pass

    def species_ring_stats(self, where="true", taxon=None):
        return [
            {"cd_ref": 60479, "anneau_ordre": 0, "nb_obs": 4, "derniere_annee": 2024, "id_rang": "ES",
             "nom_vern": "Pipistrelle commune", "lb_nom": "Pipistrellus pipistrellus", "lr_aura": "NT",
             "lr_fr_nich": "NT", "n2k": "Annexe IV"},
            {"cd_ref": 196414, "anneau_ordre": 3, "nb_obs": 2, "derniere_annee": 2022, "id_rang": "GN",
             "nom_vern": "Oreillard indéterminé", "lb_nom": "Plecotus"},
        ]

    def chiro_gites_stats(self, floutage_m=0):
        return [{"cd_ref": 196414, "anneau_ordre": 3, "type_gite": "Colonie de reproduction",
                 "nom_vern": "Oreillard indéterminé", "lb_nom": "Plecotus", "id_rang": "GN", "nb_sites": 1,
                 "effectif_max": 8, "derniere_annee": 2024}]

    def chiro_gites_resume(self, floutage_m=0):
        return {"nb_colonies": 1, "nb_gites": 0, "annee_min_colonies": 2024, "annee_max_colonies": 2024}

    def chiro_periodes(self):
        return [{"periode": "Estivage", "nb_donnees": 6, "nb_especes": 1, "nb_non_especes": 1,
                 "principales": "Pipistrelle commune (4)"}]

    def chiro_points(self, floutage_m=0):
        return [{"id": 1, "cd_nom": 60479, "lb_nom": "Pipistrellus pipistrellus", "geom": "POINT(1 1)"}]

    def chiro_gites_points(self, floutage_m=0):
        return [{"id": 1, "cd_nom": 196414, "lb_nom": "Plecotus", "geom": "POINT(1 1)"}]

    def chiro_connaissance_mailles(self):
        return []

    def hors_region(self, region):
        return {"pct_hors_region": 0, "geom": None}

    def qualite_donnees(self):
        return {"chiro": {"nb_donnees": 500, "nb_lieu_dit": 0, "nb_cachees": 0, "nb_observations": 6}}

    def chiro_gite_tampons(self, distances_m, floutage_m=0, defaut_m=1000):
        assert distances_m[196414] == 5000 and defaut_m == 1000
        return []

    def chiro_sensibilite_mailles(self, scores, score_transit_min=None):
        # période non prise en compte : pas de bonus transit
        assert scores[60479] == 3 and score_transit_min is None
        return [{"geom_maille": None, "id_area": 1, "sensi_especes": 3, "bonus_colonie": 1, "bonus_gite": 0,
                 "bonus_transit": 1, "nb_data": 6}]


class FakeSynthese:
    anneaux_km = [1.0, 6.0, 20.0]


def test_chiro_bricks_render(tmp_path, monkeypatch):
    monkeypatch.setattr(eolien_bricks, "EolienQueries", FakeChiroQueries)
    layers = {}
    monkeypatch.setattr(chiro_bricks, "_export_layer",
                        lambda ctx, rows, name, geom_col="geom": layers.setdefault(name, rows))

    ctx = ReportContext(service_name="x", id_area=1, area_name="TEST", referee="r", buffer=5, dossier=DOSSIER,
                        enabled=[], output_dirs={"tables": tmp_path, "data": tmp_path}, queries=FakeSynthese(),
                        params=dict(DOSSIER.params))
    names = ["chiro_connaissance", "chiro_vulnerabilite", "chiro_gites", "chiro_periodes", "couches_qgis_chiro"]
    keys = set()
    for name in names:
        keys |= get_brick(name).provided_keys(DOSSIER)
        ctx.results[name] = get_brick(name).run(ctx)

    texts = ctx.results["chiro_connaissance"].texts
    # 500 contacts bruts ; les comptages se font en observations (jour x lieu : nb_sessions, ou nb_obs à défaut)
    assert (texts["NB_CHIRO_DONNEES"], texts["NB_CHIRO_OBSERVATIONS"], texts["NB_CHIRO_ESPECES"],
            texts["NB_CHIRO_DONNEES_NON_ESPECE"]) == (500, 6, 1, 2)
    assert "500 données" in texts["CHIRO_CONNAISSANCE"] and "6 observations distinctes" in texts["CHIRO_CONNAISSANCE"]
    assert ctx.results["chiro_vulnerabilite"].texts["CHIRO_VULN_EOLIEN"] == "Pipistrelle commune (forte)"
    assert layers["vm_eolienne_sensi_chiro"][0]["sensibilite"] == 5  # 3 + colonie (le faux transit compte ici)
    assert layers["vm_eolienne_chiro_data"][0]["groupe"] == "pipistrelles"
    assert layers["vm_eolienne_chiro_gites"][0]["groupe"] == "oreillards"

    template = tmp_path / "template.docx"
    doc = Document()
    doc.add_paragraph("{{#chiropteres}}")
    for key in sorted(keys):
        doc.add_paragraph("{{" + key + "}}")
    doc.add_paragraph("{{/chiropteres}}")
    doc.save(template)

    assert render_report(template, ctx.results, tmp_path / "avec.docx", sections={"chiropteres"}) == set()
    assert find_placeholders(Document(tmp_path / "avec.docx")) == set()
    render_report(template, {}, tmp_path / "sans.docx", sections=set())
    assert Document(tmp_path / "sans.docx").paragraphs == [] or \
        all(not p.text for p in Document(tmp_path / "sans.docx").paragraphs)
