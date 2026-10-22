"""Cartes de dortoirs par groupe et cartes de migration par espèce dans le template Word éolien (METHODO §6, §7).

Usage : poetry run python outils/word/cartes_dortoirs_migration.py <template.docx> <sortie.docx>
- dortoirs : cartes 12.1 à 12.4 -> {{CARTES_DORTOIRS}} (une carte par groupe d'espèces) ; la carte 12 de synthèse
  reste ; phrase sur les zones de dortoir ({{DORTOIRS_RAYON}}) ajoutée au premier paragraphe ;
- migration : carte 16 (Milan) -> {{CARTES_MIGRATION_ESPECES}} (une carte par espèce) ; phrase ajoutée au premier
  paragraphe ({{MIGRATION_CARTES_SEUIL}}).
Refuse un template qui a déjà ces cartes.
"""
import sys
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph

sys.path.insert(0, str(Path(__file__).parent))
import section_docx as sd  # noqa: E402

src, dst = Path(sys.argv[1]), Path(sys.argv[2])
doc = Document(src)
W = sd.W


def para(debut=None, contient=None, style=None):
    for p in doc.paragraphs:
        t = p.text.strip()
        if (debut is None or t.startswith(debut)) and (contient is None or contient in t) \
                and (style is None or p.style.name == style):
            return p._p
    raise KeyError(debut or contient)


def remplacer_cartes(premiere, derniere, placeholder):
    """Paragraphes des cartes premiere..derniere -> un seul {{placeholder}} ; paragraphes vides intermédiaires retirés
    (jamais une fin de section)."""
    debut, fin = para(contient="{{" + premiere), para(contient="{{" + derniere)
    sd.nouveau(doc, debut.getprevious(), debut, [("{{" + placeholder + "}}", False)])
    el = debut
    while el is not None:
        suivant = el.getnext()
        texte = "".join(el.itertext()).strip()
        if el.tag == W("w:p") and el.find(W("w:pPr") + "/" + W("w:sectPr")) is None \
                and (not texte or texte.startswith("{{Carte")):
            el.getparent().remove(el)
        if el is fin:
            break
        el = suivant


assert "{{CARTES_DORTOIRS}}" not in "\n".join(p.text for p in doc.paragraphs), "cartes déjà présentes"

remplacer_cartes("Carte12.1_dortoir", "Carte12.4_dortoir", "CARTES_DORTOIRS")
Paragraph(para(debut="Sont listées dans le tableau ci-dessous"), doc._body).add_run(
    " Sur les cartes, les dortoirs d'une espèce sont représentés par des zones de {{DORTOIRS_RAYON}} de rayon autour "
    "de chaque donnée, fusionnées lorsqu'elles se touchent. Après une carte de synthèse de toutes les espèces, les "
    "espèces sont présentées par groupe (six au plus par carte), avec la même couleur que sur les cartes de "
    "nidification.")

remplacer_cartes("Carte16_migration", "Carte16_migration", "CARTES_MIGRATION_ESPECES")
Paragraph(para(debut="Au total, {{NB_MIGRATION_TOTAL}}"), doc._body).add_run(
    " Une carte par espèce présente les effectifs en migration du Milan royal et des espèces migratrices dont l'indice "
    "de vulnérabilité est de {{MIGRATION_CARTES_SEUIL}} ou plus.")
# le chapitre suivant commence sur une nouvelle page (les cartes ne remplissent plus toute la page)
Paragraph(para(debut="ANALYSE Sensibilit", style="Heading 1"), doc._body).paragraph_format.page_break_before = True

doc.save(dst)
print("dortoirs et migration : template écrit ->", dst)
