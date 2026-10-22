"""Ajoute l'indice de vulnérabilité des oiseaux (MEDDE 2015) au template Word du dossier éolien.

Usage : poetry run python outils/word/section_vulnerabilite.py <template.docx> <sortie.docx>
- chapitre « Les espèces à enjeux observées » : deux sous-titres, « Espèces à enjeux de conservation » (texte et
  tableau existants) et « Vulnérabilité à l'éolien » (nouveau, en portrait après la section paysage du tableau) ;
- phrase sur la colonne « Vulnérabilité » et les zones masquées, après la description des tableaux ;
- annexe 3 « Indice de vulnérabilité… » après l'annexe 2 ; les annexes chiroptères deviennent 4 à 6 (titres et
  renvois du texte) ;
- compteurs masqués de la numérotation des sous-titres ajustés (voir dossiers/eolien/METHODO.md §3 bis).
Refuse un template qui contient déjà la section ({{TABLE_VULNERABILITE}}).
"""
import copy
import sys
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph

sys.path.insert(0, str(Path(__file__).parent))
import section_docx as sd  # noqa: E402

src, dst = Path(sys.argv[1]), Path(sys.argv[2])
doc = Document(src)
W = sd.W


def para(debut=None, contient=None, style=None):
    for p in doc.paragraphs:
        t = p.text.strip()
        if (debut is None or t.startswith(debut)) and (contient is None or contient in t) \
                and (style is None or p.style.name == style):
            return p._p
    raise KeyError(debut or contient)


assert "{{TABLE_VULNERABILITE}}" not in "\n".join(p.text for p in doc.paragraphs), "section déjà présente"

h1_enjeux = para(debut="LES ESPECES A ENJEUX", style="Heading 1")
description = para(debut="Le tableau ci-dessous permet de visualiser chaque espèce")
table_enjeux = para(contient="{{TABLE_ESPECES_ENJEUX}}")
nidification = para(debut="Nidification", style="Heading 1")
modeles = {"titre": para(debut="Les rapaces", style="Heading 2"), "texte": description,
           "carte": table_enjeux, "legende_modele": para(debut="Tableau", contient="liste des espèces à enjeux")}
num_h2 = sd.num_heading2(doc)

# --- 1. sous-titre « Espèces à enjeux de conservation » en tête du chapitre (chapitre 5 : un compteur de plus)
h1_enjeux.addnext(sd.compteur(num_h2))
sd.nouveau(doc, h1_enjeux.getnext(), modeles["titre"], [("Espèces à enjeux de conservation", False)])

# --- 2. colonne « Vulnérabilité » et zones masquées, dans la description des tableaux
Paragraph(description, doc._body).add_run(
    " La colonne « Vulnérabilité » donne l'indice de vulnérabilité à l'éolien de l'espèce (voir la partie "
    "« Vulnérabilité à l'éolien ») ; dans tous les tableaux par zone, les zones où aucune espèce du tableau n'a "
    "été observée ne sont pas affichées.")


# --- 3. partie « Vulnérabilité à l'éolien », en portrait : juste avant le titre « Nidification »
def legende(ref, texte):
    return sd.nouveau(doc, ref, modeles["legende_modele"], [(texte, False)])


TEXTES = [
    ("titre", "Vulnérabilité à l'éolien"),
    ("texte", "En complément de la liste précédente, un indice de vulnérabilité à l'éolien hiérarchise les enjeux en "
              "combinant, pour chaque espèce, sa sensibilité à l'éolien et son enjeu de conservation. Il reprend la "
              "méthode du Protocole de suivi environnemental des parcs éoliens terrestres (MEDDE, novembre 2015) et "
              "met en avant des espèces qui ne répondent à aucun critère de la liste précédente, mais dont la "
              "vulnérabilité mérite d'être signalée."),
    ("question", "Indice de vulnérabilité = (note d'enjeu de conservation + note de sensibilité à l'éolien) / 2"),
    ("texte", "La note de sensibilité (0 à 4) est celle du protocole national : elle rapporte le nombre de cas de "
              "mortalité connus en Europe à la taille de la population européenne de l'espèce. Une espèce sans note "
              "de sensibilité (nombreux passereaux, quelques espèces absentes du référentiel) est comptée 0 : son "
              "indice repose alors sur son seul statut de conservation et peut sous-estimer sa vulnérabilité."),
    ("texte", "La note d'enjeu de conservation (1 à 5, voir la matrice ci-dessous) est la plus forte des listes "
              "rouges qui correspondent aux statuts de l'espèce sur l'aire d'étude : les listes rouges régionale "
              "et nationale des oiseaux nicheurs pour toutes les espèces ; la liste rouge nationale des oiseaux de "
              "passage seulement si l'espèce a au moins une donnée de migration ; celle des oiseaux hivernants "
              "seulement si elle a au moins une donnée entre le 1er décembre et le 31 janvier. Ces statuts sont "
              "rappelés dans la colonne « Statut sur l'aire d'étude ». Contrairement au protocole, les espèces non "
              "protégées ne sont pas ramenées à une note de 0,5 : certaines espèces chassables sont fortement "
              "menacées (Bécassine des marais, par exemple)."),
    ("legende", "Tableau : Matrice de calcul de l'indice de vulnérabilité"),
    ("bloc", "TABLE_MATRICE_VULNERABILITE"),
    ("texte", "{{TEXTE_VULNERABILITE}} Le tableau suivant les présente par indice décroissant ; l'indice de toutes "
              "les espèces observées figure en annexe 3, par ordre alphabétique."),
    ("legende", "Tableau : Espèces d'oiseaux dont l'indice de vulnérabilité est {{SEUIL_VULNERABILITE}}"),
    ("bloc", "TABLE_VULNERABILITE"),
    ("carte", "CHART_VULNERABILITE"),
    ("legende", "Figure : Nombre d'espèces d'oiseaux par valeur d'indice de vulnérabilité"),
    ("texte", "Espèces dont l'indice de vulnérabilité est de {{VULNERABILITE_PHRASES_MIN}} ou plus (éléments "
              "préremplis à vérifier et compléter : distances mesurées depuis la zone d'étude, site de nidification "
              "probable ou certain le plus proche) :"),
    ("bloc", "LISTE_ESPECES_TRES_VULNERABLES"),
    ("jaune", "{{TEXTE_COUVERTURE_CARTES}}"),
    ("jaune", "[Commentaire de l'expert : espèces les plus vulnérables, nicheuses à proximité du projet, effectifs, "
              "ancienneté des données.]"),
]
dernier = sd.inserer(doc, nidification.getprevious(), TEXTES, {**modeles, "legende": legende})

# le chapitre 5 a maintenant son compteur : un seul suffit avant « Nidification » (chapitre 6)
el, avant_nidif = nidification.getnext(), []
while el is not None and sd.est_compteur(el, num_h2):
    avant_nidif.append(el)
    el = el.getnext()
assert len(avant_nidif) == 2, f"compteurs après « Nidification » : {len(avant_nidif)}"
avant_nidif[0].getparent().remove(avant_nidif[0])

# --- 4. annexe 3 (oiseaux) ; annexes chiroptères 3-5 -> 4-6
for ancien, nouveau in (("Annexe 5", "Annexe 6"), ("Annexe 4", "Annexe 5"), ("Annexe 3", "Annexe 4"),
                        ("annexe 5", "annexe 6"), ("annexe 4", "annexe 5")):
    for p in doc.paragraphs:
        # w:t (et pas p.runs) : les titres des annexes chiroptères sont dans des liens (style « toc 1 »)
        for t in p._p.iter(W("w:t")):
            if t.text and ancien in t.text:
                t.text = t.text.replace(ancien, nouveau)
texte = "\n".join(p.text for p in doc.paragraphs)
assert "Annexe 6 : Notes de vulnérabilité des chiroptères" in texte and "annexe 6." in texte, "renumérotation"

# --- 5. carte 17 bis (indice par maille calculé à partir de la vulnérabilité), à côté de la carte 17
indice = para(debut="Pour le calcul de la sensibilité éolienne")
sd.nouveau(doc, indice, description, [(
    "Une variante de cette carte (carte 17 bis) part de l'indice de vulnérabilité le plus élevé des espèces de la "
    "maille (0,5 à 4,5, voir la partie « Vulnérabilité à l'éolien ») au lieu de leur note de sensibilité, avec les "
    "mêmes bonus : elle tient compte du statut de conservation des espèces. Les deux cartes sont proposées pour "
    "comparaison ; le choix de la méthode reste à l'appréciation du rédacteur.", False)])
carte17 = para(contient="{{Carte17_sensibilite}}")
sd.nouveau(doc, carte17, carte17, [("{{Carte17bis_vulnerabilite}}", False)])

annexe2 = para(contient="{{TABLE_AUTRES_ESPECES}}")
saut = copy.deepcopy(para(contient="{{#chiropteres}}").getnext())     # paragraphe vide avec saut de page
annexe2.addnext(saut)
titre = sd.nouveau(doc, saut, para(debut="Annexe 2", style="Heading 1"),
                   [("Annexe 3 : Indice de vulnérabilité à l'éolien des espèces d'oiseaux recensées sur le "
                     "périmètre d'étude", False)])
sd.nouveau(doc, titre, table_enjeux, [("{{TABLE_VULNERABILITE_ANNEXE}}", False)])

doc.save(dst)
print("section vulnérabilité insérée :", dst)
