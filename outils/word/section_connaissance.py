"""Ajoute la section « État des connaissances par maille de 1 km » au template Word d'un dossier.

Usage : poetry run python outils/word/section_connaissance.py <generique|eolien> <template.docx> <sortie.docx>
  generique : titre de niveau 2 à la fin de « État des connaissances », après les cartes ({{mortalite.png}}).
  eolien    : partie de niveau 1 après « Connaissance des oiseaux » (après {{Carte3_Etat_connaissance_globale}}).
Refuse un template qui contient déjà la section ({{TABLE_CONNAISSANCE_GROUPES}}).
"""
import sys
from pathlib import Path

from docx import Document

sys.path.insert(0, str(Path(__file__).parent))
import section_docx as sd  # noqa: E402
from textes_connaissance import paragraphes  # noqa: E402

profil, src, dst = sys.argv[1], Path(sys.argv[2]), Path(sys.argv[3])
doc = Document(src)


def para(debut=None, contient=None, style=None):
    for p in doc.paragraphs:
        t = p.text.strip()
        if (debut is None or t.startswith(debut)) and (contient is None or contient in t) \
                and (style is None or p.style.name == style):
            return p._p
    raise KeyError(debut or contient)


assert "{{TABLE_CONNAISSANCE_GROUPES}}" not in "\n".join(p.text for p in doc.paragraphs), "section déjà présente"

if profil == "generique":
    ancre = para(contient="{{mortalite.png}}")
    items = paragraphes(titre="État des connaissances par maille de 1 km", zone="la zone d'étude",
                        de_la_zone="de la zone d'étude", groupes="les groupes étudiés dans ce rapport",
                        porter_a_connaissance=False, legendes_cartes=False)
    modeles = {"titre": para(style="Heading 2"), "texte": para(debut="La présente synthèse provient"),
               "carte": ancre}
elif profil == "eolien":
    ancre = para(contient="{{Carte3_Etat_connaissance_globale}}")
    items = paragraphes(titre="État des connaissances par maille de 1 km", zone="la zone d'étude",
                        de_la_zone="de la zone d'étude",
                        groupes="les oiseaux et, lorsque la partie chiroptères est demandée, les chauves-souris",
                        porter_a_connaissance=False, legendes_cartes=False)
    modeles = {"titre": para(style="Heading 1", contient="OISEAUX"),
               "texte": para(debut="La présente synthèse provient"),
               "carte": ancre}
else:
    raise SystemExit(f"profil inconnu : {profil}")

# ancre : on insère après le paragraphe de la dernière carte (et ses éventuels paragraphes vides qui suivent ne
# bougent pas : la section commence juste après la carte)
sd.inserer(doc, ancre, items, modeles)
titre = ancre.getnext()

# Numérotation : dans ces templates, les titres de niveau 2 (et 3) suivent leur propre liste, dont le niveau 0
# (le chiffre du chapitre) avance grâce à des paragraphes vides masqués (« compteurs ») placés par l'auteur du
# template avant chaque chapitre qui a des sous-titres. On en ajoute ou retire pour garder des numéros justes.
num_h2 = sd.num_heading2(doc)


def compteur():
    return sd.compteur(num_h2)


def est_compteur(el):
    return sd.est_compteur(el, num_h2)


if profil == "generique":
    # « État des connaissances » est le chapitre 3 ; seul le chapitre 1 (Méthodes) a des sous-titres avant
    for _ in range(2):
        ancre.addnext(compteur())
else:
    # nouvelle partie = chapitre 4 : 3 compteurs (sous-titres « 4.x ») ; les compteurs suivants (avant
    # Nidification, chapitre 6 désormais) étaient 4 pour atteindre 5 : 2 suffisent après les nôtres
    dernier = titre
    for _ in range(3):
        dernier.addnext(compteur())
        dernier = dernier.getnext()
    suivants = []
    el = dernier.getnext()
    while el is not None and len(suivants) < 4:
        if est_compteur(el):
            suivants.append(el)
        el = el.getnext()
    assert len(suivants) == 4, f"compteurs avant Nidification : {len(suivants)}"
    for el in suivants[:2]:
        el.getparent().remove(el)

doc.save(dst)
print("section insérée :", len(items), "paragraphes ;", dst)
