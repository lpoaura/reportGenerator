"""template_plsoc.py <template_natura2000.docx> <sortie.docx> : construit le template du rapport annuel PLSOC.

Reprend la page de garde, l'historique des versions et la charte du template Natura 2000, supprime son contenu,
reconstruit le sommaire (champ TOC, mis à jour à l'ouverture avec F9) et écrit la trame du rapport annuel PLSOC
(rapports 2018-2024) : textes fixes, placeholders {{CLE}} fournis par les briques plsoc_*, paragraphes à rédiger
surlignés en jaune. Le template source n'est pas modifié.
"""

import sys

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

src, dst = sys.argv[1], sys.argv[2]
doc = Document(src)
body = doc.element.body


# ----------------------------------------------------------------------------------------------------- nettoyage
def texte(el):
    return "".join(t.text or "" for t in el.iter(qn("w:t")))


enfants = list(body)
i_sommaire = next(i for i, e in enumerate(enfants) if e.tag == qn("w:p") and texte(e).strip().upper() == "SOMMAIRE")
i_contenu = next(
    i
    for i, e in enumerate(enfants)
    if e.tag == qn("w:p")
    and e.find(qn("w:pPr")) is not None
    and e.find(qn("w:pPr")).find(qn("w:pStyle")) is not None
    and e.find(qn("w:pPr")).find(qn("w:pStyle")).get(qn("w:val")) in ("Heading1", "Titre1")
)
# ancien sommaire : entrées supprimées ; le paragraphe qui porte la fin de section garde seulement ses propriétés
for e in enfants[i_sommaire + 1:i_contenu]:
    ppr = e.find(qn("w:pPr")) if e.tag == qn("w:p") else None
    if ppr is not None and ppr.find(qn("w:sectPr")) is not None:
        for c in list(e):
            if c is not ppr:
                e.remove(c)
    else:
        body.remove(e)
# contenu Natura 2000 supprimé (la dernière propriété de section du corps est conservée)
for e in enfants[i_contenu:]:
    if e.tag != qn("w:sectPr"):
        body.remove(e)

# nouveau sommaire
som = enfants[i_sommaire]
p_toc = OxmlElement("w:p")
som.addnext(p_toc)
r = OxmlElement("w:r")
for typ, val in (
    ("begin", None),
    ("instr", ' TOC \\o "1-3" \\h \\z \\u '),
    ("separate", None),
    ("txt", None),
    ("end", None),
):
    if typ == "instr":
        e = OxmlElement("w:instrText")
        e.set(qn("xml:space"), "preserve")
        e.text = val
    elif typ == "txt":
        e = OxmlElement("w:t")
        e.text = "Sommaire : clic droit > Mettre à jour les champs (ou F9)."
    else:
        e = OxmlElement("w:fldChar")
        e.set(qn("w:fldCharType"), typ)
    r.append(e)
p_toc.append(r)

# page de garde et coordonnées (adapter à la délégation qui rédige)
for t in body.iter(qn("w:t")):
    if t.text == "[Code du site]":
        t.text = "Rapport annuel de suivi {{ANNEE}} – {{AREA_NAME}}"
REMPLACEMENTS = {
    "LPO Auvergne-Rhône-Alpes Délégation Territoriale Savoie":
        "LPO Auvergne-Rhône-Alpes – Délégation territoriale du Rhône",
    "101 Rue de Mundelsheim": "100, rue des Fougères",
    "73290 La Motte-Servolex": "69009 Lyon",
}
for p in doc.paragraphs:
    if p.text.strip() in REMPLACEMENTS:
        nouveau = REMPLACEMENTS[p.text.strip()]
        for run in p.runs[1:]:
            run.text = ""
        p.runs[0].text = nouveau
        p.runs[0].font.highlight_color = WD_COLOR_INDEX.YELLOW

# dernière section : A4 portrait (celle du template Natura 2000 est en paysage, pour ses annexes)
s, ref = doc.sections[-1], doc.sections[1]
s.orientation = WD_ORIENT.PORTRAIT
s.page_width, s.page_height = ref.page_width, ref.page_height
s.left_margin, s.right_margin, s.top_margin, s.bottom_margin = (
    ref.left_margin,
    ref.right_margin,
    ref.top_margin,
    ref.bottom_margin,
)


# ----------------------------------------------------------------------------------------------------- écriture
def style(*noms):
    existants = {st.name for st in doc.styles}
    return next(n for n in noms if n in existants)


def H(niveau, t):
    doc.add_paragraph(t, style=style(f"Heading {niveau}", f"Titre {niveau}"))


def P(t, jaune=False):
    p = doc.add_paragraph()
    r = p.add_run(t)
    if jaune:
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        r.italic = True
    return p


def X(cle):
    """Placeholder seul dans son paragraphe (tableau, image, blocs)."""
    doc.add_paragraph("{{" + cle + "}}")


def L(t):
    doc.add_paragraph(t, style=style("Caption", "Légende"))


def A(t):
    P(f"[{t}]", jaune=True)


# ----------------------------------------------------------------------------------------------------- trame
H(1, "État des données (version de travail)")
P(
    "Rapport généré automatiquement le {{DATE_DONNEES}} à partir des saisies du suivi (Mergin), de l'historique "
    "2015-2025 et des validations expertes (projet QGIS de coordination). Les chiffres, tableaux, graphiques et cartes "
    "sont produits automatiquement ; les passages surlignés en jaune sont à rédiger. Cette section est à supprimer "
    "dans la version finale."
)
X("BLOCS_ETAT_DONNEES")

H(1, "Présentation du plan de sauvegarde")
H(2, "Origines du plan")
P(
    "Le plan de sauvegarde de l'Œdicnème criard, mis en place sur le territoire de l'Est lyonnais, a pour origine la "
    "multiplication des cas de destruction de l'habitat de cette espèce sur ce territoire, nécessitant "
    "la mise en place "
    "de mesures compensatoires. Le coût et la consommation d'espace liés à ces mesures ont conduit à une prise de "
    "conscience de la nécessité d'une protection globale de l'espèce sur le territoire concerné."
)
P(
    "Dans ce contexte, la DREAL a suggéré la mise en place d'un plan local de sauvegarde. Validé par le Conseil "
    "Scientifique Régional du Patrimoine Naturel (CSRPN) en 2014, il est animé depuis 2015 par la LPO et l'APIE, en "
    "partenariat avec les collectivités du territoire."
)
A("Adapter au territoire du rapport (historique du plan, partenaires, collectivités).")
H(2, "Objectifs")
P(
    "L'objectif du plan de sauvegarde est de préserver la population d'Œdicnèmes criards tout en permettant le "
    "développement urbain de la zone. Plusieurs axes sont développés : la préservation d'une surface d'habitat "
    "suffisante, par la prise en compte de l'espèce dans les documents d'urbanisme ; le suivi et la protection de tous "
    "les couples nicheurs, avec un accent sur les couples nichant en milieu agricole ; la mise en place de mesures de "
    "compensation permettant de réduire les impacts sur les espaces agricoles tout en constituant des surfaces "
    "favorables à la reproduction de l'Œdicnème."
)
H(2, "Présentation de l'espèce")
P(
    "L'Œdicnème criard est un limicole terrestre de la famille des Burhinidés. Espèce des milieux "
    "ouverts, il recherche "
    "avant tout les milieux secs avec peu ou pas de végétation, avec une préférence pour les sols caillouteux. La plus "
    "grande partie de la population se reproduit en milieu agricole, avec une préférence pour les cultures tardives "
    "(maïs, tournesol, soja), mais l'espèce occupe aussi des milieux anthropisés (friches, bassins, carrières, "
    "toitures). La nidification débute en mars ; la couvée de deux œufs est incubée 25 à 27 jours ; les poussins, "
    "nidifuges, volent à 36-42 jours. Les pontes de remplacement sont fréquentes après un échec."
)
H(2, "Présentation de la zone d'étude")
P("{{TEXTE_ZONE}}")
X("CARTE_ZONE")
L("Carte : emprise du rapport, secteurs de suivi et mesures compensatoires.")
A("Compléter : surfaces agricoles, périmètres de protection (PENAP, ZAP), pédologie.")

H(1, "Méthodologie")
H(2, "Suivi de la population")
P(
    "Les prospections reposent sur l'observation directe des secteurs de nidification potentiels, repérés par "
    "photo-interprétation et à partir des observations des années précédentes. Les cultures semées tardivement sont "
    "ciblées en priorité. L'observation des comportements liés à la reproduction permet de classer les "
    "couples : couple "
    "observé (nicheur probable) ; couvaison observée, nid avec œufs ou présence de poussins (nicheur certain)."
)
P(
    "Lorsqu'un couple nicheur a été localisé, il est suivi régulièrement afin d'évaluer le succès de reproduction : "
    "un passage pendant la couvaison, puis tous les dix jours après l'éclosion jusqu'à l'envol des jeunes. Les "
    "observations sont saisies sur le terrain dans l'application Mergin."
)
A("Tableau du temps passé (prévisionnel / réalisé, salariés et bénévoles) à compléter.")
H(2, "Protection des nids")
P(
    "En milieu agricole, la localisation du nid est transmise à la Chambre d'agriculture, qui sollicite l'accord de "
    "l'exploitant ; le nid est ensuite balisé par deux piquets placés à 5 mètres de part et d'autre, dans le sens du "
    "travail du champ. En milieu non agricole, le degré de menace est évalué pour chaque couple ; si le "
    "site est menacé, "
    "le propriétaire est contacté pour mettre en place une protection (accès interdit, balisage)."
)
H(2, "Évaluation du succès de reproduction")
P(
    "L'issue de chaque nid est classée en succès certain, succès probable, échec probable ou échec certain selon la "
    "grille d'interprétation des observations successives (nid avec œufs ou couvaison, poussins éclos, "
    "poussins de plus "
    "de 15 jours, jeunes volants, disparition des adultes ou des poussins). En fin de saison, le succès "
    "et le nombre de "
    "jours d'exposition de chaque nid sont validés à dire d'expert ; tant qu'un nid n'est pas validé, la valeur "
    "calculée automatiquement à partir des événements saisis est utilisée et signalée comme provisoire."
)
P(
    "Le succès est aussi estimé par la méthode de Mayfield : le taux de survie journalier est égal à 1 moins le "
    "rapport entre le nombre de nids en échec et le total des jours d'exposition ; la probabilité de "
    "succès est ce taux "
    "élevé à la puissance de la durée de la nidification, soit {{MAYFIELD_DUREE}} jours de la ponte à l'envol (Cramp & "
    "Simmons, 1983)."
)
H(2, "Données utilisées")
P(
    "Ce rapport porte sur la saison {{ANNEE}} ; les séries remontent à {{ANNEE_DEBUT}}. Les données proviennent du "
    "suivi saisi dans Mergin (depuis 2026), de l'historique du suivi (2015-2025), des données Faune-France / GeoNature "
    "pour la phénologie, et des référentiels du suivi (secteurs, mesures compensatoires, PLU). Communes, secteurs et "
    "mesures compensatoires sont attribués d'après la position des couples et des nids."
)

H(1, "Résultats")
H(2, "Météorologie et conditions de réalisation")
A("Conditions météorologiques de la saison (températures, précipitations, station de référence) à rédiger.")
H(2, "Résultats généraux")
P("{{TEXTE_RESULTATS_GENERAUX}}")
X("TABLE_COMMUNES")
L("Tableau : répartition des couples et des nids par collectivité et par commune en {{ANNEE}}.")
X("CARTE_REPARTITION")
L("Carte : répartition des couples en {{ANNEE}}.")
X("CHART_POPULATION")
L("Figure : évolution de la population suivie depuis {{ANNEE_DEBUT}}.")
X("TABLE_SERIE")
L("Tableau : bilan annuel du suivi depuis {{ANNEE_DEBUT}}.")
H(2, "Milieux occupés")
P("{{TEXTE_MILIEUX}}")
X("CHART_MILIEUX_EVOLUTION")
L("Figure : évolution de la part des couples en milieu agricole.")
X("CHART_MILIEUX_CULTURES")
L("Figure : répartition des couples en milieu agricole selon la culture en {{ANNEE}}.")
X("CHART_MILIEUX_NON_AGRICOLES")
L("Figure : répartition des couples en milieu non agricole en {{ANNEE}}.")
A("Commentaire sur les milieux occupés à rédiger.")
H(2, "Secteurs de nidification")
P(
    "Le territoire est présenté par secteur de nidification. Pour chaque secteur sont précisés le nombre de couples, "
    "le statut et le succès de reproduction et, pour chaque couple, le degré de menace qui pèse sur son site de "
    "nidification : menace forte (projet d'aménagement à court terme identifié), menace modérée (site en attente de "
    "construction), menace à long terme (parcelle classée « U » ou « AU »), non menacé."
)
X("SECTEURS")
H(2, "Mise en œuvre de la protection")
P("{{TEXTE_PROTECTION}}")
X("TABLE_PROTECTION_AGRICOLE")
L("Tableau : protection des nids en milieu agricole.")
X("TABLE_PROTECTION_NON_AGRICOLE")
L("Tableau : protection des nids en milieu non agricole.")
H(2, "Phénologie et succès de reproduction")
H(3, "Phénologie de reproduction")
X("TABLE_PHENOLOGIE")
L(
    "Tableau : dates clés de la saison de reproduction depuis {{ANNEE_DEBUT}} (Hivernage : présence de "
    "l'espèce en hiver)."
)
P("{{TEXTE_PHENOLOGIE}}")
X("CHART_PHENOLOGIE")
L("Figure : nombre de couples présents par décade en {{ANNEE}} et moyenne des années précédentes.")
A("Commentaire sur la phénologie à rédiger.")
H(3, "Succès de reproduction")
X("TABLE_SUCCES")
L("Tableau : succès de reproduction des nids en {{ANNEE}}.")
P("{{TEXTE_SUCCES}}")
X("TABLE_MAYFIELD")
L("Tableau : estimation du succès de reproduction par la méthode de Mayfield en {{ANNEE}}.")
P("{{TEXTE_MAYFIELD}}")
P("{{TEXTE_CAUSES}}")
X("CHART_CAUSES")
L("Figure : causes de fin des nichées en échec en {{ANNEE}}.")
X("CHART_SUCCES_HISTORIQUE")
L("Figure : historique du succès de reproduction.")
A(
    "Commentaire sur le succès de reproduction (comparaison aux années précédentes, conditions "
    "météorologiques) à rédiger."
)
H(2, "Suivi des rassemblements postnuptiaux")
P("{{TEXTE_RASSEMBLEMENTS}}")
X("TABLE_RASSEMBLEMENTS")
L("Tableau : effectifs maximaux des rassemblements postnuptiaux en {{ANNEE}}.")
X("CHART_RASSEMBLEMENTS")
A("Description des sites de rassemblement (dates, effectifs, toitures, dérangements) à rédiger.")
H(2, "Comportements et observations remarquables")
A("Observations remarquables de la saison à rédiger.")
H(2, "Étude de la dynamique de la population : baguage")
P("{{TEXTE_BAGUAGE}}")
X("TABLE_BAGUAGE")
L("Tableau : oiseaux bagués et adultes bagués identifiés en {{ANNEE}}.")
H(2, "Bilan du suivi {{ANNEE}}")
P("{{TEXTE_BILAN}}")
A("Bilan de la saison à compléter et à commenter.")

H(1, "Prise en compte de l'Œdicnème sur le territoire")
H(2, "Nouveaux projets impactant l'Œdicnème et nouveaux adhérents")
A("Projets accompagnés (nom, porteur, état d'avancement) et nouveaux adhérents à la charte à rédiger.")
H(2, "Planification urbaine")
P("{{TEXTE_URBANISME}}")
X("TABLE_URBANISME")
L("Tableau : sites occupés depuis {{ANNEE_DEBUT}} en zone U ou AU des documents d'urbanisme.")
A("Commentaire à rédiger ; les surfaces de parcelles seront fiabilisées avec la reprise des parcelles cadastrales.")
H(2, "Suivi des mesures compensatoires")
P(
    "Les mesures compensatoires implantées sur le territoire font l'objet d'une évaluation afin d'identifier si elles "
    "permettent la nidification, ou, si ce n'est pas le cas, quels défauts de conception ou de gestion "
    "les empêchent de "
    "jouer leur rôle."
)
P("{{TEXTE_MESURES}}")
X("TABLE_MESURES")
L("Tableau : synthèse des mesures compensatoires en {{ANNEE}}.")
X("FICHES_MESURES")
H(2, "Les mesures liées aux carrières")
A(
    "Bilan des carrières (exploitant, arrêté, présence et succès de l'Œdicnème) à rédiger : aucun référentiel des "
    "carrières n'est encore en base."
)

H(1, "Conclusion générale")
A("Conclusion à rédiger.")

H(1, "Annexe : contrôle de la saisie")
P(
    "Anomalies de saisie détectées en {{ANNEE}} dans l'emprise du rapport (vue v_ctrl_anomalies) : à corriger dans "
    "Mergin avant la version définitive."
)
X("TABLE_ANOMALIES")

doc.save(dst)
print("template écrit :", dst)
