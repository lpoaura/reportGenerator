"""Convertit le rapport « Synthèse faune Combe de Savoie 2025 » en template du dossier natura2000.

Usage : poetry run python outils/word/convert_template_n2000.py <Synthese_faune_Combe_Savoie_2025.docx> <template.docx>
  (rapport de référence fourni par les collègues, gardé hors du dépôt)
Règle : chiffres, tableaux et cartes -> placeholders ; textes d'expert -> surlignés en jaune ;
partie « Résultats par secteur » supprimée.
"""
import copy
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

src, dst = sys.argv[1], sys.argv[2]
doc = Document(src)
body = doc.element.body
W = qn


def norm(t):
    return re.sub(r"\s+", " ", (t or "").replace("’", "'").replace(" ", " ")).strip()


def txt(el):
    return norm("".join(x.text or "" for x in el.iter(W("w:t"))))


def els():
    return list(body.iterchildren())


def find(start, after=None, tag="w:p"):
    """Premier paragraphe (ou tableau) dont le texte commence par start, après l'élément after."""
    seen = after is None
    for el in els():
        if not seen:
            seen = el is after
            continue
        # légendes : le numéro de figure est un champ Word (« Figure 12 : ») qu'on ignore
        texte = re.sub(r"^(Figure|Tableau) ?\d* ?:", r"\1 :", txt(el))
        if el.tag == W(tag) and texte.startswith(norm(start)):
            return el
    raise KeyError(start)


def heading(el, level):
    if el.tag != W("w:p"):
        return False
    st = Paragraph(el, doc._body).style.name
    return st == f"Heading {level}"


def has_sect(el):
    return el.tag == W("w:p") and el.find(W("w:pPr") + "/" + W("w:sectPr")) is not None


def remove(el):
    el.getparent().remove(el)


def remove_between(start_el, stop_el, keep_sect=False):
    """Supprime les éléments strictement entre start_el et stop_el (sauts de section gardés si keep_sect)."""
    cur = start_el.getnext()
    while cur is not None and cur is not stop_el:
        nxt = cur.getnext()
        if not (keep_sect and has_sect(cur)):
            remove(cur)
        cur = nxt


def set_text(p_el, parts):
    """Remplace le contenu d'un paragraphe : parts = [(texte, jaune)], mise en forme du premier run conservée."""
    first = p_el.find(".//" + W("w:r"))
    rpr = copy.deepcopy(first.find(W("w:rPr"))) if first is not None and first.find(W("w:rPr")) is not None else None
    for child in list(p_el):
        if child.tag != W("w:pPr"):
            p_el.remove(child)
    p = Paragraph(p_el, doc._body)
    for text, jaune in parts:
        r = p.add_run(text)
        if rpr is not None:
            r._r.insert(0, copy.deepcopy(rpr))
        if jaune:
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return p_el


def new_after(ref_el, model_el, parts):
    """Nouveau paragraphe (copie de la mise en forme de model_el) inséré après ref_el."""
    p = copy.deepcopy(model_el)
    sp = p.find(W("w:pPr") + "/" + W("w:sectPr"))
    if sp is not None:
        sp.getparent().remove(sp)
    set_text(p, parts)
    ref_el.addnext(p)
    return p


def jaune(p_el):
    """Tout le paragraphe en jaune (texte d'expert à reprendre)."""
    for r in p_el.iter(W("w:r")):
        rpr = r.find(W("w:rPr"))
        if rpr is None:
            rpr = OxmlElement("w:rPr")
            r.insert(0, rpr)
        for h in rpr.findall(W("w:highlight")):
            rpr.remove(h)
        h = OxmlElement("w:highlight")
        h.set(W("w:val"), "yellow")
        rpr.append(h)


def image_before(caption_el):
    """Paragraphe image qui précède une légende."""
    prev = caption_el.getprevious()
    while prev is not None and not list(prev.iter(W("w:drawing"))) and not list(prev.iter(W("w:pict"))):
        prev = prev.getprevious()
    return prev


def placeholder_image(caption_start, key, legende=None, after=None):
    cap = find(caption_start, after)
    img = image_before(cap)
    model = find("Le site Natura 2000")  # paragraphe « Normal » de référence
    p = new_after(img, model, [("{{" + key + "}}", False)])
    pPr_img = img.find(W("w:pPr"))
    if pPr_img is not None:  # alignement de l'image d'origine (centré)
        p.replace(p.find(W("w:pPr")), copy.deepcopy(pPr_img)) if p.find(W("w:pPr")) is not None else None
    remove(img)
    if legende:
        # légende : on garde le champ de numérotation, on remplace le texte après « : »
        runs = cap.findall(W("w:r"))
        done = False
        for r in runs:
            t = r.find(W("w:t"))
            if t is not None and ":" in (t.text or "") and not done:
                t.text = t.text.split(":")[0] + ": " + legende
                done = True
            elif done and t is not None:
                t.text = ""
    return p, cap


# ================================================================= couverture
for tx in body.iter(W("w:txbxContent")):
    paras = tx.findall(W("w:p"))
    t = txt(tx)
    if t.startswith("Synthèse de données du site Natura 2000"):
        set_text(paras[0], [("{{TITRE_RAPPORT}}", False)])
        for p in paras[1:]:
            if txt(p):
                set_text(p, [("[Code du site]", True)])
tables = [el for el in els() if el.tag == W("w:tbl")]
for tbl in tables[:2]:  # historique des versions, rédaction et validation
    for row in tbl.findall(W("w:tr"))[1:]:
        for p in row.iter(W("w:p")):
            jaune(p)

# ================================================================= 1. contexte
p = find("Le site Natura 2000 FR8201773")
MODEL = copy.deepcopy(p)
set_text(p, [("Le site Natura 2000 « {{NOM_SITE}} » couvre une surface de {{SURFACE_HA}} ha. ", False),
             ("[Présentation du site : code, nombre de secteurs, milieux…]", True)])
jaune(find("Afin d'actualiser l'état des lieux"))
placeholder_image("Figure : Localisation du site", "Carte01_localisation")

# ================================================================= 2. matériels et méthodes
p = find("La présente synthèse repose sur les connaissances")
set_text(p, [(txt(p) + " ", False), ("[Chiffres de la base à actualiser]", True)])
set_text(find("Pour cette synthèse, l'ensemble des observations"),
         [("Pour cette synthèse, l'ensemble des observations de faune situées au sein du site Natura 2000 et à "
           "proximité sur la période {{PERIODE_ETUDE}} ont été prises en compte. Les données antérieures n'ont pas "
           "été considérées.", False)])
set_text(find("Un tampon de 1 kilomètre"), [("{{TEXTE_RAYONS}}", False)])
placeholder_image("Figure : Répartition de la pression d'observation", "Carte02_pression_observation")

# ================================================================= 3. résultats
jaune(find("Le site Natura 2000 peut être découpé"))
set_text(find("La synthèse des connaissances se fera d'abord"),
         [("La synthèse des connaissances porte sur l'ensemble du site Natura 2000. ", False),
          ("[Présentation des sous-secteurs à rédiger si besoin]", True)])

# 3.1 ensemble du site
CAP_MODEL = find("Figure : Répartition de la pression d'observation")


def legende_apres(ref_el, texte):
    """Légende « Figure <n> : texte » (copie d'une légende du modèle, numéro = champ Word) après ref_el."""
    c = copy.deepcopy(CAP_MODEL)
    ref_el.addnext(c)
    for r in c.findall(W("w:r")):
        t = r.find(W("w:t"))
        if t is not None and ":" in (t.text or ""):
            t.text = t.text.split(":")[0] + ": " + texte
        elif t is not None and t.text and not t.text.strip().startswith("Figure"):
            t.text = ""
    return c


chiffres = find("Sur la période 2015-2025")
set_text(chiffres, [("{{TEXTE_CHIFFRES_SITE}}", False)])
# graphiques natifs Word aux chiffres figés (répartition des espèces / des données par groupe) -> graphiques générés
for el in [e for e in els() if list(e.iter("{http://schemas.openxmlformats.org/drawingml/2006/chart}chart"))]:
    remove(el)
g = new_after(chiffres, MODEL, [("{{CHART_ESPECES_GROUPES}}", False)])
g = legende_apres(g, "Nombre d'espèces par groupe taxonomique dans le site")
g = new_after(g, MODEL, [("{{CHART_DONNEES_GROUPES}}", False)])
legende_apres(g, "Nombre de données par groupe taxonomique dans le site")
jaune(find("Parmi les autres taxons, les papillons de jour"))
nb = find("NB : il est important de rappeler")
h3_model = next(el for el in els() if heading(el, 3))

# 3.1.1 état des connaissances par groupe (porter à connaissance) : texte commun aux dossiers
sys.path.insert(0, str(Path(__file__).parent))
import section_docx  # noqa: E402
from textes_connaissance import paragraphes  # noqa: E402

n = section_docx.inserer(doc, nb, paragraphes(), {
    "titre": h3_model, "texte": MODEL, "carte": MODEL,
    "legende": lambda ref, texte: legende_apres(ref, texte)})

# 3.1.2 espèces Natura 2000
h = new_after(n, h3_model, [("Espèces Natura 2000", False)])
p = find("La présente synthèse se focalisera principalement")
h.addnext(p)  # le paragraphe d'introduction passe sous le nouveau titre
set_text(p, [("La présente synthèse se focalisera principalement sur les espèces déterminantes pour le programme "
              "Natura 2000, soit les espèces inscrites à l'annexe I de la directive Oiseaux ou à l'annexe II de la "
              "directive Habitats Faune Flore, soit {{NB_ESPECES_N2000}} espèces listées dans le tableau suivant. "
              "La colonne « Présence » indique si "
              "l'espèce est connue dans le site ou seulement à proximité (dans le rayon de prise en compte de son "
              "groupe, cf. Matériels et méthodes).", False)])
tbl = find("Groupe", p, tag="w:tbl")
t_p = new_after(tbl, MODEL, [("{{TABLE_ESPECES_N2000}}", False)])
remove(tbl)
new_after(t_p, MODEL, [("{{TEXTE_COMPLEXES}}", False)])
p = find("A celles-ci on peut ajouter")
set_text(p, [("À celles-ci on peut ajouter d'autres espèces déterminantes à grand rayon d'action, recensées "
              "uniquement à proximité du site : {{LISTE_ESPECES_GRAND_RAYON}}.", False)])
for start in ("Loup gris", "Lynx boréal"):
    el = find(start, p)
    if has_sect(el):  # porte un saut de section : on le garde, vidé et sans puce
        set_text(el, [("", False)])
        for num in el.iter(W("w:numPr")):
            num.getparent().remove(num)
    else:
        remove(el)

# 3.1.x espèces : textes, photos et cartes par espèce -> atlas (titres par groupe créés à la génération)
oiseaux = next(el for el in els() if heading(el, 3) and txt(el).startswith("Oiseaux nicheurs"))
chiro = next(el for el in els() if heading(el, 3) and txt(el).startswith("Chiroptères"))
intro = new_after(oiseaux.getprevious(), MODEL, [
    ("Les espèces Natura 2000 sont présentées ci-dessous par groupe ; chaque espèce fait l'objet d'une carte de "
     "localisation des données. Au total, {{NB_OISEAUX_NICHEURS_N2000}} espèces d'oiseaux avec un statut de nicheur "
     "et inscrites à l'annexe I de la directive Oiseaux sont connues sur le site et à proximité. ", False),
    ("[Classement des oiseaux par cortège (zones humides, forestiers, rupestres, milieux ouverts) et texte de "
     "présentation de chaque espèce à rédiger, avec une photo]", True)])
atlas = new_after(intro, MODEL, [("{{ATLAS_ESPECES_N2000}}", False)])
remove_between(atlas, chiro)

# chiroptères
set_text(chiro, [("Chiroptères", False)])
set_text(find("Si aucune donnée de chiroptère", chiro),
         [("{{NB_CHIRO_N2000}} espèces de chiroptères inscrites à l'annexe II de la directive Habitats sont connues "
           "dans le site et à proximité. ", False), ("[Commentaire à rédiger]", True)])
tbl = find("Nom vernaculaire", chiro, tag="w:tbl")
new_after(tbl, MODEL, [("{{TABLE_CHIRO_N2000}}", False)])
remove(tbl)
jaune(find("L'absence d'observations est plus le signe", chiro))
items = [find(s, chiro) for s in ("Les observations de toutes les espèces Natura 2000",
                                  "La localisation de toutes les gîtes connus sans reproduction",
                                  "La localisation de toutes les colonies de reproduction connues",
                                  "Les données des colonies de reproduction d'espèce Natura 2000",
                                  "Les localisations de tous les gîtes connus, par espèce")]
textes = ["Les observations des espèces Natura 2000, que ce soit des observations opportunistes ou de gîtes.",
          "La localisation de tous les gîtes et colonies de reproduction connus, pour tous les chiroptères, avec "
          "l'effectif maximal observé.",
          "Les colonies de reproduction connues, par espèce (toutes espèces).",
          "Les gîtes sans reproduction des espèces Natura 2000.",
          "Les gîtes et colonies connus, par espèce Natura 2000."]
for el, t in zip(items, textes):
    set_text(el, [(t, False)])
for cap, key, leg in [("Figure : Observation de chiroptères N2000", "Carte04_chiro_presence_n2000",
                       "Présence des chiroptères Natura 2000"),
                      ("Figure : : Gîtes et colonies connus, toutes espèces", "Carte05_chiro_gites_colonies",
                       "Gîtes et colonies connus, toutes espèces"),
                      ("Figure : Présence de colonies de reproduction", "Carte06_chiro_colonies_especes",
                       "Colonies de reproduction connues par espèce"),
                      ("Figure : Gîte d'espèces Natura 2000 sans reproduction", "Carte07_chiro_gites_n2000",
                       "Gîtes sans reproduction des espèces Natura 2000")]:
    _, last_cap = placeholder_image(cap, key, leg, after=chiro)
# cartes par espèce (6 dans le modèle) -> atlas des gîtes
secteurs = next(el for el in els() if heading(el, 2) and txt(el).startswith("Résultats par secteur"))
atlas_chiro = new_after(last_cap, MODEL, [("{{ATLAS_CHIRO_GITES}}", False)])
# section portrait de référence : la dernière avant les cartes chiroptères
avant = els()[:els().index(atlas_chiro)]
portrait = [e for e in avant if has_sect(e)][-1].find(W("w:pPr") + "/" + W("w:sectPr"))
cur = atlas_chiro.getnext()
while cur is not None and cur is not secteurs:
    nxt = cur.getnext()
    if has_sect(cur):
        # section A3 paysage des cartes chiroptères -> A4 portrait comme le reste (cartes en portrait)
        sp = cur.find(W("w:pPr") + "/" + W("w:sectPr"))
        for tag in ("w:pgSz", "w:pgMar"):
            old, new = sp.find(W(tag)), portrait.find(W(tag))
            sp.replace(old, copy.deepcopy(new))
        for child in list(cur):
            if child.tag != W("w:pPr"):
                cur.remove(child)
    else:
        remove(cur)
    cur = nxt

# ================================================================= 3.2 secteurs : supprimé
discussion = next(el for el in els() if heading(el, 1) and txt(el).startswith("Discussions"))
cur = secteurs
while cur is not None and cur is not discussion:
    nxt = cur.getnext()
    remove(cur)  # y compris le saut de section paysage des secteurs
    cur = nxt

# ================================================================= 4. discussion
# (le paragraphe de l'annexe 1 contient aussi une zone de texte « Annexe 1 (bis) » : recherche par contenu)
annexe = next(e for e in els()[els().index(discussion):] if "Liste des oiseaux recensés sur le site" in txt(e))
fin_disc = next(e for e in els()[els().index(discussion):els().index(annexe)] if has_sect(e))
d = new_after(discussion, MODEL, [("[Discussion et propositions d'actions à rédiger : principaux enjeux de "
                                   "conservation du site (milieux, espèces), propositions de gestion et "
                                   "d'amélioration des connaissances.]", True)])
remove_between(d, fin_disc)

# ================================================================= annexes : un tableau par groupe
a = new_after(fin_disc, MODEL, [("{{TABLES_ANNEXES}}", False)])
cur = a.getnext()
while cur is not None and cur.tag != W("w:sectPr"):
    nxt = cur.getnext()
    remove(cur)
    cur = nxt

# sommaire et numéros de figures : mis à jour par Word à l'ouverture (les anciens titres et signets n'existent plus)
settings = doc.settings.element
if settings.find(W("w:updateFields")) is None:
    uf = OxmlElement("w:updateFields")
    uf.set(W("w:val"), "true")
    # ordre imposé par le schéma : avant hdrShapeDefaults / footnotePr / endnotePr / compat / ...
    suivants = ("hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids", "mathPr",
                "attachedSchema", "themeFontLang", "clrSchemeMapping", "doNotIncludeSubdocsInStats",
                "doNotAutoCompressPictures", "forceUpgrade", "captions", "readModeInkLockDown", "smartTagType",
                "schemaLibrary", "shapeDefaults", "doNotEmbedSmartTags", "decimalSymbol", "listSeparator")
    apres = next((c for c in settings if c.tag.split("}")[1] in suivants), None)
    if apres is not None:
        apres.addprevious(uf)
    else:
        settings.append(uf)

# images supprimées du texte (photos, anciennes cartes) : on retire aussi les fichiers du package
xml = body.xml
orphelins = [rid for rid, rel in doc.part.rels.items()
             if ("image" in rel.reltype or rel.reltype.endswith("/chart")) and f'"{rid}"' not in xml]
for rid in orphelins:
    del doc.part.rels[rid]
print(len(orphelins), "images orphelines retirées")

doc.save(dst)
print("template écrit :", dst)
