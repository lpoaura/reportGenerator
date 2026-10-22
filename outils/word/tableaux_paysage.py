"""Met des tableaux larges du template éolien sur des pages A4 paysage (la suite repasse en portrait).

Usage : poetry run python outils/word/tableaux_paysage.py <template.docx> <sortie.docx> [PLACEHOLDER ...]
  Par défaut : TABLE_VULNERABILITE (indice de vulnérabilité des oiseaux) et TABLE_CHIRO_VULNERABILITE.
Pour chaque tableau : fin de la section portrait avant sa légende (« Tableau : … » juste au-dessus), section A4
paysage jusqu'au tableau inclus (même mise en page que les tableaux des ardéidés et des passereaux), puis la section
d'origine reprend. Les fins de section restent à l'intérieur d'une section conditionnelle ({{#chiropteres}}) : si elle
est retirée, la mise en page est inchangée. Refuse un tableau déjà en paysage.
"""
import sys
from pathlib import Path

from docx import Document

sys.path.insert(0, str(Path(__file__).parent))
import section_docx as sd  # noqa: E402

src, dst = Path(sys.argv[1]), Path(sys.argv[2])
placeholders = sys.argv[3:] or ["TABLE_VULNERABILITE", "TABLE_CHIRO_VULNERABILITE"]
doc = Document(src)
W = sd.W


def a4_paysage():
    """sectPr A4 paysage existant du template (tableaux des ardéidés)."""
    for el in doc.element.body.iter(W("w:sectPr")):
        pg = el.find(W("w:pgSz"))
        if pg is not None and pg.get(W("w:orient")) == "landscape" and pg.get(W("w:w")) == "16838":
            return el
    raise KeyError("aucune section A4 paysage dans le template")


paysage = a4_paysage()
for nom in placeholders:
    table = next((p._p for p in doc.paragraphs if p.text.strip() == "{{" + nom + "}}"), None)
    assert table is not None, f"{{{{{nom}}}}} absent du template"
    courante = sd.sectpr_apres(table)
    assert courante.find(W("w:pgSz")).get(W("w:orient")) != "landscape", f"{nom} : déjà en paysage"
    legende = table.getprevious()
    assert "".join(legende.itertext()).strip().startswith("Tableau"), f"{nom} : légende « Tableau » attendue"
    sd.fin_de_section(legende, courante, avant=True)   # tout ce qui précède reste en portrait
    sd.fin_de_section(table, paysage)                   # légende + tableau en A4 paysage
    print(f"{nom} : A4 paysage")

doc.save(dst)
print("template écrit ->", dst)
