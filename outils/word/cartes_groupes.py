"""Cartes de nidification par groupe d'espèces dans le template Word du dossier éolien (METHODO.md §5 bis).

Usage : poetry run python outils/word/cartes_groupes.py <template.docx> <sortie.docx>
- chapitre Nidification : paragraphe de lecture des cartes (après le titre) ;
- cartes 6 à 9 -> {{CARTES_NIDIF_RAPACES}}, 10 à 11.1 -> {{CARTES_NIDIF_GRANDS_VOILIERS}},
  11.2 -> {{CARTES_NIDIF_PASSEREAUX}} (une image par carte du plan, insérée après le rendu QGIS) ;
- sous-partie « Autres espèces nicheuses à enjeux » après les passereaux : texte, tableau {{TABLE_AUTRES_NICHEURS}}
  (section A4 paysage, comme les ardéidés et les passereaux), puis {{CARTES_NIDIF_AUTRES}} (portrait).
Refuse un template qui a déjà les cartes par groupe ({{CARTES_NIDIF_RAPACES}}).
"""
import sys
from pathlib import Path

from docx import Document

sys.path.insert(0, str(Path(__file__).parent))
import section_docx as sd  # noqa: E402

src, dst = Path(sys.argv[1]), Path(sys.argv[2])
doc = Document(src)
W = sd.W


def para(debut=None, contient=None, style=None):
    for p in doc.paragraphs:
        t = p.text.strip()
        if (debut is None or t.startswith(debut)) and (contient is None or contient in t) \
                and (style is None or p.style.name == style):
            return p._p
    raise KeyError(debut or contient)


def a_sectpr(el):
    return el.find(W("w:pPr") + "/" + W("w:sectPr")) is not None


def remplacer_cartes(premiere, derniere, placeholder):
    """Paragraphes des cartes premiere..derniere -> un seul paragraphe {{placeholder}} ; les paragraphes vides
    intermédiaires (espacements, sauts de page) sont retirés, jamais ceux qui portent une fin de section."""
    debut, fin = para(contient="{{" + premiere), para(contient="{{" + derniere)
    sd.nouveau(doc, debut.getprevious(), debut, [("{{" + placeholder + "}}", False)])
    el = debut
    while el is not None:
        suivant = el.getnext()
        texte = "".join(el.itertext()).strip()
        if el.tag == W("w:p") and not a_sectpr(el) and (not texte or texte.startswith("{{Carte")):
            el.getparent().remove(el)
        if el is fin:
            break
        el = suivant


assert "{{CARTES_NIDIF_RAPACES}}" not in "\n".join(p.text for p in doc.paragraphs), "cartes par groupe déjà présentes"
texte = para(debut="Au total, {{NB_RAPACE_NICHEUSE}}")

# --- 1. lecture des cartes, en tête du chapitre (avant le premier sous-titre)
rapaces = para(debut="Les rapaces", style="Heading 2")
sd.nouveau(doc, rapaces.getprevious(), texte, [(
    "Les cartes présentent, groupe d'espèces par groupe d'espèces, les données de nidification probable ou certaine "
    "(possible, probable ou certaine pour les passereaux) de toutes les espèces des tableaux, à raison de six espèces "
    "au plus par carte : un groupe plus grand est découpé (par exemple hiboux et chouettes), et les espèces "
    "qu'aucun groupe ne reconnaît figurent sur une carte « autres espèces ». Les données d'une espèce sont regroupées "
    "par maille de 100 m (point au centre des données) ; un point grossi et cerclé de noir signale un site occupé au "
    "moins trois années différentes. Lorsque des espèces se superposent (colonie, lieu-dit), leurs points sont "
    "disposés en cercle autour de leur position. Chaque espèce garde la même couleur sur toutes les cartes du "
    "rapport. Les sites des espèces les plus sensibles sont floutés (Cigogne noire : maille de 5 km).", False)])

# --- 2. cartes figées -> cartes par groupe
remplacer_cartes("Carte6_nidif", "Carte9_nidif", "CARTES_NIDIF_RAPACES")
remplacer_cartes("Carte10_nidif", "Carte11.1_nidif", "CARTES_NIDIF_GRANDS_VOILIERS")
remplacer_cartes("Carte11.2_nidif", "Carte11.2_nidif", "CARTES_NIDIF_PASSEREAUX")

# --- 3. sous-partie « Autres espèces nicheuses à enjeux » (même mise en page que les passereaux)
cartes_passereaux = para(contient="{{CARTES_NIDIF_PASSEREAUX}}")
paysage = sd.sectpr_avant(para(contient="{{TABLE_PASSEREAUX_NICHEURS}}").getnext())
portrait = sd.sectpr_avant(para(debut="Les dortoirs", style="Heading 1").getprevious())
modeles = {"titre": para(debut="Les passereaux", style="Heading 2"), "texte": texte, "carte": cartes_passereaux,
           "legende_modele": para(debut="Tableau", contient="Passereaux nicheurs", style="Standard")}


def legende(ref, txt):
    return sd.nouveau(doc, ref, modeles["legende_modele"], [(txt, False)])


fin_portrait = sd.fin_de_section(cartes_passereaux, portrait)
sd.inserer(doc, fin_portrait, [
    ("titre", "Autres espèces nicheuses à enjeux"),
    ("texte", "Sont retenues ici les autres espèces nicheuses (nicheur probable ou certain) à enjeux, qui ne sont ni "
              "des rapaces, ni des grands voiliers, ni des passereaux : espèces menacées (CR, EN ou VU sur la liste "
              "rouge régionale ou nationale des oiseaux nicheurs) ou dont l'indice de vulnérabilité est "
              "{{SEUIL_VULNERABILITE}}. On en compte {{NB_AUTRES_NICHEURS}} dans un rayon de {{RAYON_MAX_KM}} km. "
              "Elles sont cartographiées par groupe (canards, limicoles, râles…) ; celles qu'aucun groupe ne "
              "reconnaît figurent sur la carte « autres espèces »."),
    ("jaune", "[DESCRIPTION DES ENJEUX SUR LES ESPECES PRESENTES]"),
    ("legende", "Tableau : Autres espèces nicheuses à enjeux ({{RAYON_MAX_KM}} km)"),
    ("bloc", "TABLE_AUTRES_NICHEURS"),
], {**modeles, "legende": legende})
table_autres = para(contient="{{TABLE_AUTRES_NICHEURS}}")
fin_paysage = sd.fin_de_section(table_autres, paysage)
sd.nouveau(doc, fin_paysage, cartes_passereaux, [("{{CARTES_NIDIF_AUTRES}}", False)])

doc.save(dst)
print("cartes par groupe : template écrit ->", dst)
