"""Ajoute la sous-partie « Les passereaux » (chapitre Nidification) au template Word du dossier éolien.

Usage : poetry run python outils/word/section_passereaux.py <template.docx> <sortie.docx>
Après « Les Ardéidés, Ciconiidés et Laridés » (dernière carte {{Carte11.1_nidif_Hérons}}), avant « Les dortoirs » :
titre de niveau 2 (même chapitre : pas de compteur à ajuster), texte fixe, tableau {{TABLE_PASSEREAUX_NICHEURS}},
commentaire à rédiger et carte {{Carte11.2_nidif_passereaux}}.
Refuse un template qui contient déjà la sous-partie.
"""

import sys
from pathlib import Path

from docx import Document

sys.path.insert(0, str(Path(__file__).parent))
import section_docx as sd  # noqa: E402

src, dst = Path(sys.argv[1]), Path(sys.argv[2])
doc = Document(src)


def para(debut=None, contient=None, style=None):
    for p in doc.paragraphs:
        t = p.text.strip()
        if (debut is None or t.startswith(debut)) and (contient is None or contient in t) \
                and (style is None or p.style.name == style):
            return p._p
    raise KeyError(debut or contient)


assert "{{TABLE_PASSEREAUX_NICHEURS}}" not in "\n".join(p.text for p in doc.paragraphs), "sous-partie déjà présente"

carte_herons = para(contient="{{Carte11.1_nidif_")
modeles = {"titre": para(debut="Les rapaces", style="Heading 2"),
           "texte": para(debut="Au total, {{NB_ARDEIDES_NICHEUSE}}"),
           "carte": carte_herons,
           "legende_modele": para(debut="Tableau", contient="ardéidés", style="Standard")}


def legende(ref, texte):
    return sd.nouveau(doc, ref, modeles["legende_modele"], [(texte, False)])


TEXTES = [
    ("titre", "Les passereaux"),
    ("texte", "Les passereaux sont moins exposés aux collisions que les rapaces ou les grands voiliers, mais leurs "
              "territoires de reproduction et d'alimentation peuvent être directement touchés par l'implantation "
              "d'un parc éolien (défrichement, ouverture de pistes et de plateformes, dérangement pendant les "
              "travaux). Leur vulnérabilité tient le plus souvent à leur statut de conservation. Elle mérite d'autant "
              "plus d'être signalée que ces petites espèces sont difficilement détectées lors des suivis de "
              "mortalité."),
    ("texte", "Sont retenus les passereaux nicheurs (nicheur possible, probable ou certain) dans un rayon de "
              "{{RAYON_MAX_KM}} km dont l'indice de vulnérabilité est {{SEUIL_VULNERABILITE}} : "
              "{{NB_PASSEREAUX_VULNERABLES}} espèce(s) parmi les {{NB_PASSEREAUX_NICHEURS}} passereaux nicheurs "
              "observés. La carte présente leurs données de nidification, avec une couleur par espèce."),
    ("jaune", "[DESCRIPTION DES ENJEUX SUR LES ESPECES PRESENTES : milieux occupés, secteurs, proximité du projet]"),
    ("legende", "Tableau : Passereaux nicheurs dont l'indice de vulnérabilité est {{SEUIL_VULNERABILITE}} "
                "({{RAYON_MAX_KM}} km)"),
    ("bloc", "TABLE_PASSEREAUX_NICHEURS"),
    ("carte", "Carte11.2_nidif_passereaux"),
]
sd.inserer(doc, carte_herons, TEXTES, {**modeles, "legende": legende})

# Mise en page comme les ardéidés : texte et tableau sur une section A4 paysage, carte en portrait.
# Une section s'arrête au paragraphe qui porte son w:sectPr : un paragraphe vide « fin de portrait » avant le titre,
# un paragraphe vide « fin de paysage » après le tableau (sectPr copiés de ceux qui encadrent les ardéidés).
W = sd.W
paysage = sd.sectpr_avant(carte_herons)                      # section des ardéidés (A4 paysage)
portrait = sd.sectpr_avant(para(debut="Les dortoirs", style="Heading 1").getprevious())  # cartes (A4 portrait)
assert paysage.find(W("w:pgSz")).get(W("w:orient")) == "landscape", "section des ardéidés : paysage attendu"
sd.fin_de_section(para(debut="Les passereaux", style="Heading 2"), portrait, avant=True)
sd.fin_de_section(para(contient="{{TABLE_PASSEREAUX_NICHEURS}}"), paysage)

doc.save(dst)
print("sous-partie passereaux insérée :", dst)
